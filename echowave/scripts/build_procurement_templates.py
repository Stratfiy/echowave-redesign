"""Build the standard Indian procurement formats as Word templates.

Run from the repo root:

    python -m scripts.build_procurement_templates

Writes ``api/services/documents/standard/{purchase_order,rfq,work_order,
comparative_statement,award_letter}.docx``. The files are committed; this
script is how they are changed, so a layout fix is a reviewed diff here
rather than an edit in Word nobody can see.

Placeholders are ``{{field_name}}``; a table row holding ``{{items.<col>}}``
repeats once per line item (see api/services/documents/templates.py). Every
money figure is computed by money.py and placed here already formatted.
Neutral colours, one sans-serif face, A4 with 15 mm margins.
"""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

OUT = (
    Path(__file__).resolve().parent.parent
    / "api"
    / "services"
    / "documents"
    / "standard"
)

INK = RGBColor(0x22, 0x26, 0x2B)
MUTED = RGBColor(0x5F, 0x66, 0x6D)
RULE = "9AA1A9"
FILL = "EEF0F2"
FONT = "Arial"


# ---------------------------------------------------------------------------
# Small helpers over python-docx


def _document(*, landscape: bool = False) -> Document:
    document = Document()
    section = document.sections[0]
    section.page_height, section.page_width = Mm(297), Mm(210)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_height, section.page_width = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(15))
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.space_before = Pt(0)
    return document


def _para(
    container,
    text: str = "",
    *,
    bold=False,
    size=None,
    color=None,
    align=None,
    space_after=None,
    space_before=None,
):
    paragraph = container.add_paragraph()
    if text:
        run = paragraph.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    if align is not None:
        paragraph.alignment = align
    if space_after is not None:
        paragraph.paragraph_format.space_after = Pt(space_after)
    if space_before is not None:
        paragraph.paragraph_format.space_before = Pt(space_before)
    return paragraph


def _label_value(container, label: str, value: str, *, bold_value=False):
    paragraph = container.add_paragraph()
    label_run = paragraph.add_run(f"{label}: ")
    label_run.font.color.rgb = MUTED
    value_run = paragraph.add_run(value)
    value_run.bold = bold_value
    paragraph.paragraph_format.space_after = Pt(1)
    return paragraph


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _borders(table, *, color: str = RULE, inside: bool = True, size: int = 4) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    edges = ["top", "left", "bottom", "right"] + (
        ["insideH", "insideV"] if inside else []
    )
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        if edge in edges:
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), str(size))
            element.set(qn("w:color"), color)
        else:
            element.set(qn("w:val"), "nil")
        borders.append(element)
    properties.append(borders)


def _no_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    properties.append(borders)


def _widths(table, widths_mm: list[float]) -> None:
    """Fixed column widths, in the grid as well as on each cell: LibreOffice
    (and so the PDF) reads the grid, Word reads the cells."""
    table.autofit = False
    for column, width in zip(table.columns, widths_mm):
        column.width = Mm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_mm):
            cell.width = Mm(width)
    properties = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    properties.append(layout)
    total = OxmlElement("w:tblW")
    total.set(qn("w:w"), str(int(sum(widths_mm) * 56.7)))
    total.set(qn("w:type"), "dxa")
    for existing in properties.findall(qn("w:tblW")):
        properties.remove(existing)
    properties.append(total)


def _cell_text(cell, text: str, *, bold=False, align=None, size=None, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    return paragraph


def _cell_lines(cell, lines: list[tuple[str, str | None]]) -> None:
    """Label/value lines in one cell; a label of None is a plain line."""
    cell.text = ""
    first = True
    for label, value in lines:
        paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        paragraph.paragraph_format.space_after = Pt(1)
        if label is None:
            run = paragraph.add_run(value or "")
            run.bold = True
            continue
        label_run = paragraph.add_run(f"{label}: ")
        label_run.font.color.rgb = MUTED
        paragraph.add_run(value or "")


def _letterhead(document, title: str) -> None:
    """Buyer in the page header, the document number in the footer."""
    section = document.sections[0]
    header = section.header
    first = header.paragraphs[0]
    run = first.add_run("{{buyer_name}}")
    run.bold = True
    run.font.size = Pt(13)
    lines = header.add_paragraph()
    muted = lines.add_run("{{buyer_address}}")
    muted.font.size = Pt(8.5)
    muted.font.color.rgb = MUTED
    gstin = header.add_paragraph()
    gstin_run = gstin.add_run("GSTIN: {{buyer_gstin}}")
    gstin_run.font.size = Pt(8.5)
    gstin_run.font.color.rgb = MUTED
    rule = header.add_paragraph()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), RULE)
    border.append(bottom)
    rule._p.get_or_add_pPr().append(border)

    footer = section.footer.paragraphs[0]
    footer_run = footer.add_run(f"{title} {{{{document_number}}}}  ·  Page ")
    footer_run.font.size = Pt(7.5)
    footer_run.font.color.rgb = MUTED
    _page_number(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _para(
        document,
        title.upper(),
        bold=True,
        size=15,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=2,
        space_after=8,
    )


def _page_number(paragraph) -> None:
    run = paragraph.add_run()
    run.font.size = Pt(7.5)
    run.font.color.rgb = MUTED
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)


def _parties(
    document, *, vendor_label: str, right: list[tuple[str, str]], pan: bool = True
) -> None:
    table = document.add_table(rows=1, cols=2)
    _borders(table)
    left = table.rows[0].cells[0]
    left.text = ""
    head = left.paragraphs[0]
    head_run = head.add_run(vendor_label)
    head_run.font.color.rgb = MUTED
    head_run.font.size = Pt(8)
    name = left.add_paragraph()
    name.add_run("{{vendor_name}}").bold = True
    left.add_paragraph("{{vendor_address}}")
    _label_value(left, "GSTIN", "{{vendor_gstin}}")
    if pan:
        _label_value(left, "PAN", "{{vendor_pan}}")
    for paragraph in left.paragraphs:
        paragraph.paragraph_format.space_after = Pt(1)
    _cell_lines(table.rows[0].cells[1], [(label, value) for label, value in right])
    _widths(table, [100, 80])
    _para(document, space_after=4)


def _items_table(document, columns: list[tuple[str, str, float, str]]) -> None:
    """columns: (heading, placeholder or literal, width mm, align l|c|r)."""
    table = document.add_table(rows=2, cols=len(columns))
    _borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    aligns = {
        "l": WD_ALIGN_PARAGRAPH.LEFT,
        "c": WD_ALIGN_PARAGRAPH.CENTER,
        "r": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    for cell, (heading, _, _, align) in zip(table.rows[0].cells, columns):
        _cell_text(cell, heading, bold=True, size=8, align=aligns[align])
        _shade(cell, FILL)
    for cell, (_, value, _, align) in zip(table.rows[1].cells, columns):
        _cell_text(cell, value, size=8, align=aligns[align])
    _widths(table, [width for _, _, width, _ in columns])
    # Repeat the heading row on every page of a long order.
    row_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    row_properties.append(repeat)


def _totals(document, rows: list[tuple[str, str]], *, width_label=45, width_value=35):
    table = document.add_table(rows=len(rows), cols=2)
    _borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    for index, (label, value) in enumerate(rows):
        last = index == len(rows) - 1
        _cell_text(table.rows[index].cells[0], label, bold=last, size=8.5)
        _cell_text(
            table.rows[index].cells[1],
            value,
            bold=last,
            size=8.5,
            align=WD_ALIGN_PARAGRAPH.RIGHT,
        )
        if last:
            _shade(table.rows[index].cells[0], FILL)
            _shade(table.rows[index].cells[1], FILL)
    _widths(table, [width_label, width_value])


def _section(document, heading: str) -> None:
    _para(document, heading, bold=True, size=10, space_before=8, space_after=2)


def _signatory(document, *, acceptance: str | None = None) -> None:
    _para(document, space_after=6)
    table = document.add_table(rows=1, cols=2)
    _no_borders(table)
    left, right = table.rows[0].cells
    if acceptance:
        _cell_lines(left, [(None, acceptance)])
        left.add_paragraph()
        left.add_paragraph()
        left.add_paragraph("Signature, name and seal of the vendor")
        left.add_paragraph("Date:")
    right.text = ""
    first = right.paragraphs[0]
    first.add_run("For {{buyer_name}}").bold = True
    first.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for _ in range(3):
        right.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.RIGHT
    name = right.add_paragraph()
    name.add_run("{{signatory_name}}").bold = True
    name.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    title = right.add_paragraph("{{signatory_designation}}")
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = right.add_paragraph()
    label_run = label.add_run("Authorised Signatory")
    label_run.font.color.rgb = MUTED
    label_run.font.size = Pt(8)
    label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _widths(table, [90, 90])


def _terms_block(document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=len(rows), cols=2)
    _borders(table)
    for index, (label, value) in enumerate(rows):
        _cell_text(table.rows[index].cells[0], label, bold=True, size=8.5)
        _shade(table.rows[index].cells[0], FILL)
        _cell_text(table.rows[index].cells[1], value, size=8.5)
    _widths(table, [45, 135])


PRICED_COLUMNS = [
    ("Sl", "{{items.sl}}", 7, "c"),
    ("Description", "{{items.description}}", 38, "l"),
    ("HSN/SAC", "{{items.hsn_sac}}", 13, "c"),
    ("Qty", "{{items.qty}}", 10, "r"),
    ("Unit", "{{items.unit}}", 10, "c"),
    ("Rate (₹)", "{{items.rate}}", 19, "r"),
    ("Disc%", "{{items.discount}}", 9, "r"),
    ("Taxable (₹)", "{{items.taxable_value}}", 22, "r"),
    ("GST%", "{{items.gst_rate}}", 9, "r"),
    ("GST (₹)", "{{items.gst_amount}}", 20, "r"),
    ("Amount (₹)", "{{items.amount}}", 23, "r"),
]

TAX_TOTALS = [
    ("Taxable value (₹)", "{{subtotal}}"),
    ("CGST (₹)", "{{cgst}}"),
    ("SGST (₹)", "{{sgst}}"),
    ("IGST (₹)", "{{igst}}"),
    ("Total (₹)", "{{total}}"),
]


# ---------------------------------------------------------------------------
# The five formats


def purchase_order() -> Document:
    document = _document()
    _letterhead(document, "Purchase Order")
    _parties(
        document,
        vendor_label="SUPPLIER",
        right=[
            ("PO No.", "{{document_number}}"),
            ("PO Date", "{{document_date}}"),
            ("Quotation ref.", "{{reference}}"),
            ("Delivery by", "{{delivery_date}}"),
        ],
    )
    _para(
        document,
        "Please supply the following on the terms below, quoting this PO number "
        "on your invoice and delivery challan.",
        space_after=4,
    )
    _items_table(document, PRICED_COLUMNS)
    _para(document, space_after=2)
    _totals(document, TAX_TOTALS)
    _para(document, space_after=2)
    _label_value(document, "Amount in words", "{{amount_in_words}}", bold_value=True)
    _section(document, "Terms")
    _terms_block(
        document,
        [
            ("Deliver to", "{{delivery_address}}"),
            ("Delivery by", "{{delivery_date}}"),
            ("Payment terms", "{{payment_terms}}"),
            ("Validity of this PO", "{{validity}}"),
            ("Terms and conditions", "{{terms_and_conditions}}"),
        ],
    )
    _signatory(
        document,
        acceptance="Accepted on the terms above",
    )
    return document


def work_order() -> Document:
    document = _document()
    _letterhead(document, "Work Order")
    _parties(
        document,
        vendor_label="CONTRACTOR",
        right=[
            ("WO No.", "{{document_number}}"),
            ("WO Date", "{{document_date}}"),
            ("Quotation ref.", "{{reference}}"),
            ("Complete by", "{{delivery_date}}"),
        ],
    )
    _section(document, "Scope of work")
    _para(document, "{{scope_of_work}}", space_after=4)
    _items_table(document, PRICED_COLUMNS)
    _para(document, space_after=2)
    _totals(document, TAX_TOTALS)
    _para(document, space_after=2)
    _label_value(document, "Amount in words", "{{amount_in_words}}", bold_value=True)
    _section(document, "Terms")
    _terms_block(
        document,
        [
            ("Site / place of work", "{{delivery_address}}"),
            ("Completion by", "{{delivery_date}}"),
            ("Payment terms", "{{payment_terms}}"),
            ("Validity of this order", "{{validity}}"),
            ("Terms and conditions", "{{terms_and_conditions}}"),
        ],
    )
    _signatory(document, acceptance="Accepted on the terms above")
    return document


def rfq() -> Document:
    document = _document()
    _letterhead(document, "Request for Quotation")
    _parties(
        document,
        vendor_label="TO",
        right=[
            ("RFQ No.", "{{document_number}}"),
            ("RFQ Date", "{{document_date}}"),
            ("Reference", "{{reference}}"),
            ("Quote by", "{{quotation_due_date}}"),
        ],
        pan=False,
    )
    _para(
        document,
        "We invite your best quotation for the items below. Please quote the "
        "rate per unit, the GST rate and the HSN/SAC code for each line, and "
        "send it to us by {{quotation_due_date}}.",
        space_after=4,
    )
    _items_table(
        document,
        [
            ("Sl", "{{items.sl}}", 8, "c"),
            ("Description / specification", "{{items.description}}", 70, "l"),
            ("HSN/SAC", "{{items.hsn_sac}}", 18, "c"),
            ("Qty", "{{items.qty}}", 14, "r"),
            ("Unit", "{{items.unit}}", 14, "c"),
            ("Rate (₹)", "", 22, "r"),
            ("GST %", "", 14, "r"),
            ("Remarks", "", 20, "l"),
        ],
    )
    _section(document, "Terms")
    _terms_block(
        document,
        [
            ("Deliver to", "{{delivery_address}}"),
            ("Required by", "{{delivery_date}}"),
            ("Payment terms", "{{payment_terms}}"),
            ("Quote valid for", "{{validity}}"),
            ("Terms and conditions", "{{terms_and_conditions}}"),
        ],
    )
    _signatory(document)
    return document


def comparative_statement() -> Document:
    document = _document()
    _letterhead(document, "Comparative Statement")
    table = document.add_table(rows=1, cols=2)
    _borders(table)
    _cell_lines(
        table.rows[0].cells[0],
        [("Subject", "{{subject}}"), ("RFQ reference", "{{reference}}")],
    )
    _cell_lines(
        table.rows[0].cells[1],
        [("CS No.", "{{document_number}}"), ("Date", "{{document_date}}")],
    )
    _widths(table, [110, 70])
    _para(document, space_after=4)
    _items_table(
        document,
        [
            ("Sl", "{{items.sl}}", 8, "c"),
            ("Vendor", "{{items.vendor_name}}", 38, "l"),
            ("GSTIN", "{{items.vendor_gstin}}", 30, "c"),
            ("Quoted value (₹)", "{{items.quoted_total}}", 24, "r"),
            ("Delivery", "{{items.delivery_period}}", 18, "c"),
            ("Payment", "{{items.payment_terms}}", 20, "c"),
            ("Rank", "{{items.rank}}", 12, "c"),
            ("Remarks", "{{items.remarks}}", 30, "l"),
        ],
    )
    _section(document, "Recommendation")
    _para(document, "{{recommendation}}", space_after=3)
    _terms_block(
        document,
        [
            ("Recommended vendor", "{{vendor_name}}"),
            ("Recommended value (₹)", "{{total}}"),
            ("In words", "{{amount_in_words}}"),
        ],
    )
    _para(
        document,
        "The item-wise comparison of rates, landed cost and ranking is attached "
        "as a spreadsheet.",
        color=MUTED,
        size=8,
        space_before=4,
    )
    _signatory(document)
    return document


def award_letter() -> Document:
    document = _document()
    _letterhead(document, "Letter of Award")
    _parties(
        document,
        vendor_label="TO",
        right=[
            ("Ref. No.", "{{document_number}}"),
            ("Date", "{{document_date}}"),
            ("Your quotation", "{{reference}}"),
        ],
    )
    _para(document, "Subject: {{subject}}", bold=True, space_after=6)
    _para(document, "Dear Sir / Madam,", space_after=4)
    _para(
        document,
        "With reference to your quotation {{reference}}, we are pleased to "
        "award you the contract for {{subject}} for a total value of "
        "₹ {{total}} ({{amount_in_words}}), inclusive of applicable taxes, on "
        "the terms below. A purchase / work order with the detailed schedule "
        "follows.",
        space_after=6,
    )
    _terms_block(
        document,
        [
            ("Contract value (₹)", "{{total}}"),
            ("Delivery / completion", "{{delivery_date}}"),
            ("Payment terms", "{{payment_terms}}"),
            ("Please accept within", "{{validity}}"),
            ("Terms and conditions", "{{terms_and_conditions}}"),
        ],
    )
    _para(
        document,
        "Please sign and return a copy of this letter as your acceptance.",
        space_before=6,
    )
    _signatory(document, acceptance="Accepted")
    return document


BUILDERS = {
    "purchase_order": purchase_order,
    "rfq": rfq,
    "work_order": work_order,
    "comparative_statement": comparative_statement,
    "award_letter": award_letter,
}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for name, build in BUILDERS.items():
        path = OUT / f"{name}.docx"
        document = build()
        document.core_properties.created = datetime(2026, 9, 1, tzinfo=UTC)
        document.core_properties.modified = datetime(2026, 9, 1, tzinfo=UTC)
        document.core_properties.last_modified_by = "Decibyl"
        document.core_properties.author = "Decibyl"
        document.core_properties.title = name.replace("_", " ").title()
        document.save(path)
        print(f"wrote {path.relative_to(OUT.parent.parent.parent.parent)}")


if __name__ == "__main__":
    main()
