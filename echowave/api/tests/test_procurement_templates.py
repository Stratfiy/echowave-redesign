"""Filling a Word template: ``{{field}}`` wherever Word put it.

Word splits text into runs whenever formatting, spell-check or an edit
history says so, and ``{{vendor_name}}`` typed into a real template is as
likely to arrive as ``{{vend`` + ``or_na`` + ``me}}`` as in one piece. A
filler that only looks inside runs fills the test document and misses the
customer's -- so the tests build the split the way Word does.
"""

import io

import docx
import pytest

from api.services.documents import templates


def _bytes(document) -> bytes:
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _split_paragraph(paragraph, pieces, bold_first=True):
    for index, piece in enumerate(pieces):
        run = paragraph.add_run(piece)
        if index == 0 and bold_first:
            run.bold = True


def _all_text(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    for section in document.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(parts)


@pytest.fixture
def template() -> bytes:
    document = docx.Document()
    header = document.sections[0].header.paragraphs[0]
    header.text = "{{buyer_name}} -- Purchase Order"
    document.sections[0].footer.paragraphs[0].text = "PO {{document_number}}"
    _split_paragraph(
        document.add_paragraph(),
        ["To: {{vend", "or_na", "me}}, GSTIN {", "{vendor_gstin}}"],
    )
    document.add_paragraph("Date: {{ document_date }}")
    table = document.add_table(rows=3, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Sl", "Description", "Amount"]):
        cell.text = text
    for cell, text in zip(
        table.rows[1].cells,
        ["{{items.sl}}", "{{items.description}}", "{{items.amount}}"],
    ):
        cell.text = text
    table.rows[2].cells[1].text = "Total"
    table.rows[2].cells[2].text = "{{total}}"
    return _bytes(document)


class TestInspect:
    def test_it_finds_every_field_including_split_ones(self, template):
        found = templates.inspect(template)
        assert found["fields"] == [
            "buyer_name",
            "document_number",
            "vendor_name",
            "vendor_gstin",
            "document_date",
            "total",
        ]
        assert found["item_columns"] == ["sl", "description", "amount"]

    def test_an_unreadable_file_says_so(self):
        with pytest.raises(templates.TemplateError):
            templates.inspect(b"not a docx")


class TestFill:
    def test_split_run_placeholders_are_filled_and_keep_the_first_runs_format(
        self, template
    ):
        filled = templates.fill(
            template,
            {
                "buyer_name": "Acme",
                "vendor_name": "Bharat Steels",
                "vendor_gstin": "29AAGCB7383J1Z4",
                "document_date": "26-09-2026",
                "document_number": "PO/26-27/0001",
                "total": "1,180.00",
            },
            [],
        )
        document = docx.Document(io.BytesIO(filled))
        paragraph = document.paragraphs[0]
        assert paragraph.text == "To: Bharat Steels, GSTIN 29AAGCB7383J1Z4"
        assert paragraph.runs[0].bold is True
        assert document.paragraphs[1].text == "Date: 26-09-2026"
        assert "{{" not in _all_text(filled)

    def test_headers_and_footers_are_filled(self, template):
        filled = templates.fill(
            template, {"buyer_name": "Acme", "document_number": "PO/26-27/0007"}, []
        )
        document = docx.Document(io.BytesIO(filled))
        assert document.sections[0].header.paragraphs[0].text == (
            "Acme -- Purchase Order"
        )
        assert document.sections[0].footer.paragraphs[0].text == "PO PO/26-27/0007"

    def test_a_repeating_row_is_cloned_per_item_in_order(self, template):
        items = [
            {"sl": 1, "description": "Cement", "amount": "100.00"},
            {"sl": 2, "description": "Sand", "amount": "50.00"},
            {"sl": 3, "description": "Steel", "amount": "900.00"},
        ]
        filled = templates.fill(template, {"total": "1,050.00"}, items)
        table = docx.Document(io.BytesIO(filled)).tables[0]
        rows = [[c.text for c in r.cells] for r in table.rows]
        assert rows == [
            ["Sl", "Description", "Amount"],
            ["1", "Cement", "100.00"],
            ["2", "Sand", "50.00"],
            ["3", "Steel", "900.00"],
            ["", "Total", "1,050.00"],
        ]

    def test_no_items_removes_the_repeating_row(self, template):
        filled = templates.fill(template, {}, [])
        table = docx.Document(io.BytesIO(filled)).tables[0]
        assert len(table.rows) == 2

    def test_a_multiline_value_keeps_its_lines(self, template):
        filled = templates.fill(template, {"vendor_name": "Line one\nLine two"}, [])
        paragraph = docx.Document(io.BytesIO(filled)).paragraphs[0]
        assert "Line one\nLine two" in paragraph.text

    def test_an_unknown_placeholder_is_left_visible(self, template):
        # The tool refuses to draft with a field missing; if one slips
        # through, a visible {{field}} is caught by whoever reads the draft,
        # where a blank would not be.
        filled = templates.fill(template, {"buyer_name": "Acme"}, [])
        assert "{{vendor_gstin}}" in _all_text(filled)

    def test_text_around_placeholders_is_untouched(self):
        document = docx.Document()
        document.add_paragraph("No placeholders here.")
        filled = templates.fill(_bytes(document), {"x": "y"}, [])
        assert docx.Document(io.BytesIO(filled)).paragraphs[0].text == (
            "No placeholders here."
        )


class TestFromText:
    def test_plain_text_becomes_a_template_with_an_items_table(self):
        text = (
            "Purchase Order {{document_number}}\n"
            "Vendor: {{vendor_name}}\n"
            "{{items.description}} | {{items.qty}} | {{items.amount}}\n"
            "Total: {{total}}"
        )
        data = templates.docx_from_text(text)
        found = templates.inspect(data)
        assert found["fields"] == ["document_number", "vendor_name", "total"]
        assert found["item_columns"] == ["description", "qty", "amount"]
        filled = templates.fill(
            data,
            {"document_number": "PO/1", "vendor_name": "V", "total": "3"},
            [{"description": "a", "qty": "1", "amount": "1"}] * 3,
        )
        assert len(docx.Document(io.BytesIO(filled)).tables[0].rows) == 4
