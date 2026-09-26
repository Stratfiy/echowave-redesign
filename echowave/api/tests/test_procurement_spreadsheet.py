"""Spreadsheets an agent hands over, opened the way a person would open them.

Each test reads the workbook back with openpyxl and checks what somebody
looking at it in Excel sees: a bold, frozen heading row, rupee formats on the
money columns, the formulas that make it a working sheet rather than a
picture of one, and the lowest bid on each line marked.
"""

import io
from decimal import Decimal

import openpyxl
import pytest

from api.services.documents import spreadsheet


def _open(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data))


class TestBuildWorkbook:
    def test_heading_frozen_formats_formulas_and_widths(self):
        data = spreadsheet.build_workbook(
            "Cement quotes",
            [
                {
                    "name": "Quotes",
                    "columns": ["Vendor", "Qty", "Rate (₹)", "Amount"],
                    "rows": [
                        ["Ultratech dealer", 100, 385.5, 38550],
                        ["ACC dealer", 100, 379, 37900],
                    ],
                    "number_formats": {"Qty": "0"},
                    "formulas": [
                        {"cell": "D4", "formula": "SUM(D2:D3)"},
                        {"cell": "C4", "formula": "=MIN(C2:C3)"},
                    ],
                    "column_widths": {"A": 30},
                }
            ],
        )
        book = _open(data)
        sheet = book["Quotes"]
        assert [c.value for c in sheet[1]] == ["Vendor", "Qty", "Rate (₹)", "Amount"]
        assert all(c.font.bold for c in sheet[1])
        assert sheet.freeze_panes == "A2"
        assert "₹" in sheet["C2"].number_format
        assert "₹" in sheet["D2"].number_format
        assert sheet["B2"].number_format == "0"
        assert sheet["D4"].value == "=SUM(D2:D3)"
        assert sheet["C4"].value == "=MIN(C2:C3)"
        assert "₹" in sheet["D4"].number_format
        assert sheet.column_dimensions["A"].width == 30
        # Auto width: wide enough for the longest value, never absurd.
        assert 8 <= sheet.column_dimensions["B"].width <= 60
        assert book.properties.title == "Cement quotes"

    def test_money_columns_can_be_named(self):
        data = spreadsheet.build_workbook(
            "x",
            [
                {
                    "name": "S",
                    "columns": [{"name": "Landed", "type": "money"}, "Note"],
                    "rows": [[12.5, "a"]],
                }
            ],
        )
        sheet = _open(data)["S"]
        assert "₹" in sheet["A2"].number_format
        assert "₹" not in sheet["B2"].number_format

    def test_a_formula_cell_must_be_a_cell(self):
        with pytest.raises(spreadsheet.SpreadsheetError):
            spreadsheet.build_workbook(
                "x",
                [
                    {
                        "name": "S",
                        "columns": ["a"],
                        "rows": [],
                        "formulas": [{"cell": "not a cell", "formula": "SUM(A1)"}],
                    }
                ],
            )

    def test_sheet_names_are_made_legal(self):
        data = spreadsheet.build_workbook(
            "x", [{"name": "Q1/Q2: [bids]?", "columns": ["a"], "rows": [[1]]}]
        )
        assert _open(data).sheetnames == ["Q1-Q2- -bids--"]


VENDORS = [
    {
        "name": "Sri Balaji Traders",
        "rates": [380, 62, 54000, 410],
        "freight": 5,
        "gst_rate": 18,
    },
    {
        "name": "Karnataka Build Mart",
        "rates": [372, 65, 53500, 405],
        "freight": 8,
        "gst_rate": 18,
    },
    {
        "name": "Deccan Supplies",
        "rates": [390, 60, 55000, None],
        "freight": 0,
        "gst_rate": 18,
    },
]
ITEMS = [
    {"description": "OPC 53 cement", "qty": 200, "unit": "Bag"},
    {"description": "M-sand", "qty": 30, "unit": "Cft"},
    {"description": "TMT 12mm Fe550D", "qty": 4, "unit": "MT"},
    {"description": "Binding wire", "qty": 50, "unit": "Kg"},
]


class TestCostBidAnalysis:
    def test_landed_cost_totals_and_ranking(self):
        result = spreadsheet.cost_bid_analysis(VENDORS, ITEMS)
        by_vendor = {row["vendor"]: row for row in result.ranking}
        # Sri Balaji: (380+5)*1.18=454.30 x200 = 90,860.00; (62+5)*1.18=79.06 x30
        # = 2,371.80; (54000+5)*1.18=63,725.90 x4 = 2,54,903.60;
        # (410+5)*1.18=489.70 x50 = 24,485.00. Total 3,72,620.40.
        assert by_vendor["Sri Balaji Traders"]["total"] == Decimal("372620.40")
        # Deccan did not quote binding wire, so it ranks after the complete bids.
        assert [r["vendor"] for r in result.ranking][-1] == "Deccan Supplies"
        assert result.ranking[0]["rank"] == "L1"
        assert [r["rank"] for r in result.ranking] == ["L1", "L2", "L3"]
        assert result.ranking[-1]["complete"] is False

    def test_lowest_per_item(self):
        result = spreadsheet.cost_bid_analysis(VENDORS, ITEMS)
        # M-sand: Balaji 79.06, Build Mart 86.14, Deccan 70.80.
        assert result.lowest[1] == "Deccan Supplies"
        # Cement: Balaji 454.30, Build Mart 448.40, Deccan 460.20.
        assert result.lowest[0] == "Karnataka Build Mart"

    def test_the_workbook(self):
        result = spreadsheet.cost_bid_analysis(VENDORS, ITEMS)
        book = _open(spreadsheet.build_workbook("Cost-bid analysis", result.sheets))
        assert book.sheetnames == ["Summary", "Comparison"]
        comparison = book["Comparison"]
        assert comparison.freeze_panes == "A2"
        header = [c.value for c in comparison[1]]
        assert header[:4] == ["Sl", "Item", "Qty", "Unit"]
        assert "Karnataka Build Mart landed/unit (₹)" in header
        # A SUM row under each vendor's totals and a MIN per item.
        formulas = [
            c.value
            for row in comparison.iter_rows()
            for c in row
            if isinstance(c.value, str) and c.value.startswith("=")
        ]
        assert any(f.startswith("=SUM(") for f in formulas)
        assert any(f.startswith("=MIN(") for f in formulas)
        # The lowest landed cost on the cement line is highlighted.
        landed_col = header.index("Karnataka Build Mart landed/unit (₹)") + 1
        cell = comparison.cell(row=2, column=landed_col)
        assert cell.fill.fgColor.rgb.endswith(spreadsheet.LOWEST_FILL)
        other = comparison.cell(
            row=2, column=header.index("Sri Balaji Traders landed/unit (₹)") + 1
        )
        assert not str(other.fill.fgColor.rgb).endswith(spreadsheet.LOWEST_FILL)
        assert "₹" in cell.number_format

        summary = book["Summary"]
        assert [c.value for c in summary[2]][:2] == ["L1", result.ranking[0]["vendor"]]
        assert summary.freeze_panes == "A2"

    def test_a_vendor_without_rates_is_refused(self):
        with pytest.raises(spreadsheet.SpreadsheetError):
            spreadsheet.cost_bid_analysis([{"name": "X", "rates": [1]}], ITEMS)

    def test_a_negative_rate_is_refused(self):
        with pytest.raises(spreadsheet.SpreadsheetError):
            spreadsheet.cost_bid_analysis(
                [{"name": "X", "rates": [-1, 1, 1, 1]}], ITEMS
            )
