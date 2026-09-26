"""The document tools as an agent calls them.

The two paths that matter most: a draft with anything missing produces no
file and no number, and asks every question at once; a complete draft is
numbered, filled, converted, stored under this workspace's keys, entered in
the register awaiting approval, and handed over on the thread with its files
-- and is not sent.
"""

import io
from datetime import UTC, datetime
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import docx
import openpyxl
import pytest
from sqlalchemy import select

from api import constants
from api.db.models import OrganizationModel, ProcurementDocumentModel
from api.enums import AgentEventKind
from api.services.documents import formats, money, templates, tools

BUYER = money.with_checksum("29AABCT1332L1Z")
VENDOR = money.with_checksum("29AAGCB7383J1Z")


class FakeStorage:
    def __init__(self):
        self.files: dict[str, bytes] = {}

    async def acreate_file_from_bytes(self, key, data):
        self.files[key] = data
        return True

    async def aget_signed_url(self, key, expiration=3600, **_):
        return f"https://files.test/{key}?ttl={expiration}"

    async def aread_bytes(self, key, max_bytes):
        return self.files.get(key)


@pytest.fixture
def flag_on(monkeypatch):
    monkeypatch.setattr(constants, "PROCUREMENT_DOCS_2026_09_ENABLED", True)


@pytest.fixture
def storage(monkeypatch):
    from api.services import storage as storage_module

    fake = FakeStorage()
    monkeypatch.setattr(storage_module, "storage_fs", fake)
    return fake


@pytest.fixture
def converter():
    with patch(
        "api.services.documents.convert.docx_to_pdf",
        new=AsyncMock(return_value=b"%PDF-1.7 test"),
    ) as mocked:
        yield mocked


@pytest.fixture
def timeline():
    with patch(
        "api.services.workflow.agent_timeline.record", new=AsyncMock()
    ) as mocked:
        yield mocked


@pytest.fixture
def charges():
    with patch(
        "api.services.billing.events.charge_in_own_session",
        new=AsyncMock(return_value=1),
    ) as mocked:
        yield mocked


async def _org(db_session) -> int:
    async with db_session.async_session() as session:
        org = OrganizationModel(
            provider_id=f"proc-tools-{datetime.now(UTC).timestamp()}",
            quota_decibyl_tokens=0,
        )
        session.add(org)
        await session.flush()
        return org.id


def _complete_values() -> dict:
    return {
        "buyer_name": "Shreeram Constructions Pvt Ltd",
        "buyer_address": "42, 2nd Main, Peenya Industrial Area, Bengaluru 560058",
        "buyer_gstin": BUYER,
        "vendor_name": "Bharat Building Supplies",
        "vendor_address": "Plot 7, KIADB, Hosur Road, Bengaluru 560100",
        "vendor_gstin": VENDOR,
        "vendor_pan": VENDOR[2:12],
        "reference": "Q-2291 dated 20-09-2026",
        "delivery_address": "Site: Whitefield, Bengaluru",
        "delivery_date": "2026-10-10",
        "payment_terms": "30 days from receipt of invoice",
        "validity": "30 days",
        "terms_and_conditions": "standard",
        "signatory_name": "R. Srinivas",
        "signatory_designation": "Purchase Manager",
    }


def _items() -> list[dict]:
    return [
        {
            "description": "OPC 53 cement",
            "hsn_sac": "2523",
            "qty": 200,
            "unit": "Bag",
            "rate": 380,
            "gst_rate": 28,
        },
        {
            "description": "TMT 12mm",
            "hsn_sac": "7214",
            "qty": 2,
            "unit": "MT",
            "rate": "54000",
            "discount": 2,
            "gst_rate": 18,
        },
    ]


class TestPrepare:
    def test_everything_missing_is_asked_at_once(self):
        inspected = templates.inspect(
            formats.FORMATS["purchase_order"].template_bytes()
        )
        prepared = tools.prepare("purchase_order", inspected, {"vendor_name": "X"}, [])
        asked = {m["field"] for m in prepared.missing}
        assert "buyer_gstin" in asked and "payment_terms" in asked and "items" in asked
        # Worked out, never asked.
        for derived in (
            "document_number",
            "document_date",
            "total",
            "cgst",
            "amount_in_words",
        ):
            assert derived not in asked
        assert "vendor_name" not in asked
        assert all(
            m["question"].endswith("?") or m["question"].endswith(".")
            for m in prepared.missing
        )

    def test_an_item_gap_is_asked_about_that_item(self):
        inspected = templates.inspect(
            formats.FORMATS["purchase_order"].template_bytes()
        )
        items = _items()
        del items[1]["hsn_sac"]
        prepared = tools.prepare("purchase_order", inspected, _complete_values(), items)
        assert prepared.missing == [
            {
                "field": "items[2].hsn_sac",
                "question": "For item 2 (TMT 12mm), what is the HSN or SAC code?",
            }
        ]

    def test_a_bad_gstin_and_a_negative_quantity_are_errors(self):
        inspected = templates.inspect(
            formats.FORMATS["purchase_order"].template_bytes()
        )
        values = {**_complete_values(), "vendor_gstin": "29AAGCB7383J1ZZ"}
        items = _items()
        items[0]["qty"] = -5
        prepared = tools.prepare("purchase_order", inspected, values, items)
        assert any("vendor_gstin" in e for e in prepared.errors)
        assert any("item 1 qty" in e for e in prepared.errors)

    def test_complete_values_are_printed_indian_style(self):
        inspected = templates.inspect(
            formats.FORMATS["purchase_order"].template_bytes()
        )
        prepared = tools.prepare(
            "purchase_order", inspected, _complete_values(), _items()
        )
        assert prepared.missing == [] and prepared.errors == []
        # Cement 76,000 at 28% = 21,280; TMT 1,08,000 less 2% = 1,05,840 at 18%
        # = 19,051.20. Same state: split in halves.
        assert prepared.values["subtotal"] == "1,81,840.00"
        assert prepared.values["cgst"] == "20,165.60"
        assert prepared.values["sgst"] == "20,165.60"
        assert prepared.values["igst"] == "0.00"
        assert prepared.values["total"] == "2,22,171.20"
        assert prepared.values["amount_in_words"] == (
            "Rupees Two Lakh Twenty-Two Thousand One Hundred Seventy-One and "
            "Twenty Paise Only"
        )
        assert prepared.values["terms_and_conditions"] == formats.STANDARD_TERMS
        assert prepared.items[1]["taxable_value"] == "1,05,840.00"
        assert prepared.items[1]["sl"] == "2"

    def test_an_award_letter_takes_its_total(self):
        inspected = templates.inspect(formats.FORMATS["award_letter"].template_bytes())
        prepared = tools.prepare(
            "award_letter",
            inspected,
            {
                **_complete_values(),
                "subject": "Supply of cement",
                "total": "2,22,171.20",
            },
            [],
        )
        assert prepared.missing == []
        assert prepared.values["amount_in_words"].startswith("Rupees Two Lakh")


@pytest.mark.asyncio
class TestDraftDocument:
    async def test_missing_fields_draft_nothing_and_use_no_number(
        self, flag_on, db_session, storage, converter, timeline, charges
    ):
        org = await _org(db_session)
        result = await tools.run(
            tools.DRAFT,
            organization_id=org,
            arguments={"kind": "purchase_order", "values": {"vendor_name": "Bharat"}},
            ref_id="t:1",
        )
        assert result["status"] == "missing"
        assert {m["field"] for m in result["missing"]} >= {"buyer_gstin", "items"}
        assert "ONE message" in result["note"]
        assert storage.files == {}
        converter.assert_not_awaited()
        timeline.assert_not_awaited()
        charges.assert_not_awaited()
        async with db_session.async_session() as session:
            rows = (
                await session.scalars(
                    select(ProcurementDocumentModel).where(
                        ProcurementDocumentModel.organization_id == org
                    )
                )
            ).all()
        assert rows == []

    async def test_a_complete_draft_is_numbered_filed_and_handed_over(
        self, flag_on, db_session, storage, converter, timeline, charges
    ):
        org = await _org(db_session)
        result = await tools.run(
            tools.DRAFT,
            organization_id=org,
            arguments={
                "kind": "purchase_order",
                "template": "purchase_order",
                "values": _complete_values(),
                "items": _items(),
                "counterparty_email": "sales@bharat.example",
            },
            ref_id="t:2",
            workflow_id=None,
        )
        assert result["status"] == "drafted", result
        number = result["number"]
        assert number.startswith("PO/") and number.endswith("/0001")
        stem = number.replace("/", "-")
        assert [f["filename"] for f in result["files"]] == [
            f"{stem}.pdf",
            f"{stem}.docx",
        ]
        assert all(
            f["url"].startswith("https://files.test/procurement/")
            for f in result["files"]
        )
        assert "ttl=604800" in result["files"][0]["url"]
        assert "NOT been sent" in result["note"]

        # Stored under this workspace's own prefix.
        keys = sorted(storage.files)
        assert all(
            k.startswith(f"procurement/{org}/{result['register_id']}/") for k in keys
        )
        docx_bytes = next(v for k, v in storage.files.items() if k.endswith(".docx"))
        text = "\n".join(
            p.text for p in docx.Document(io.BytesIO(docx_bytes)).paragraphs
        )
        cells = "\n".join(
            c.text
            for t in docx.Document(io.BytesIO(docx_bytes)).tables
            for r in t.rows
            for c in r.cells
        )
        assert "{{" not in text + cells
        assert number in cells
        assert "2,22,171.20" in cells
        assert "OPC 53 cement" in cells
        converter.assert_awaited_once()

        async with db_session.async_session() as session:
            row = await session.scalar(
                select(ProcurementDocumentModel).where(
                    ProcurementDocumentModel.id == result["register_id"]
                )
            )
        assert row.organization_id == org
        assert row.status == "awaiting_approval"
        assert row.amount_paise == 22217120
        assert row.counterparty_gstin == VENDOR
        assert row.due_date.isoformat() == "2026-10-10"
        assert row.data["counterparty_email"] == "sales@bharat.example"

        event = timeline.await_args.kwargs
        assert event["kind"] == AgentEventKind.DELIVERABLE.value
        assert event["organization_id"] == org
        attachments = event["payload"]["attachments"]
        assert [a["filename"] for a in attachments] == [f"{stem}.pdf", f"{stem}.docx"]
        assert all(set(a) >= {"document_uuid", "filename"} for a in attachments)
        assert attachments[0]["register_id"] == result["register_id"]
        assert attachments[0]["file"] == "pdf"
        charges.assert_awaited_once()
        assert charges.await_args.kwargs["ref_id"] == "t:2"

    async def test_an_invalid_gstin_drafts_nothing(
        self, flag_on, db_session, storage, converter, timeline, charges
    ):
        org = await _org(db_session)
        result = await tools.run(
            tools.DRAFT,
            organization_id=org,
            arguments={
                "kind": "purchase_order",
                "values": {**_complete_values(), "buyer_gstin": "29AABCT1332L1ZX"},
                "items": _items(),
            },
            ref_id="t:3",
        )
        assert result["status"] == "invalid"
        assert storage.files == {}

    async def test_the_second_draft_takes_the_next_number(
        self, flag_on, db_session, storage, converter, timeline, charges
    ):
        org = await _org(db_session)
        numbers = []
        for i in range(2):
            result = await tools.run(
                tools.DRAFT,
                organization_id=org,
                arguments={
                    "kind": "purchase_order",
                    "values": _complete_values(),
                    "items": _items(),
                },
                ref_id=f"t:n{i}",
            )
            numbers.append(result["number"])
        assert [n[-4:] for n in numbers] == ["0001", "0002"]


@pytest.mark.asyncio
class TestOtherTools:
    async def test_list_template_fields(self, flag_on, charges):
        result = await tools.run(
            tools.LIST_FIELDS,
            organization_id=1,
            arguments={"template": "rfq"},
            ref_id="t:f",
        )
        assert result["status"] == "success"
        assert "quotation_due_date" in result["fields"]
        assert result["item_columns"] == ["sl", "description", "hsn_sac", "qty", "unit"]

    async def test_cost_bid_spreadsheet_is_stored_and_handed_over(
        self, flag_on, db_session, storage, timeline, charges
    ):
        org = await _org(db_session)
        result = await tools.run(
            tools.SPREADSHEET,
            organization_id=org,
            arguments={
                "title": "Cement and steel bids",
                "kind": "cost_bid_analysis",
                "items": [{"description": "Cement", "qty": 10, "unit": "Bag"}],
                "vendors": [
                    {"name": "A", "rates": [380]},
                    {"name": "B", "rates": [370]},
                ],
            },
            ref_id="t:x",
        )
        assert result["status"] == "success", result
        assert result["ranking"][0]["vendor"] == "B"
        assert result["number"].startswith("CS/")
        key, data = next(iter(storage.files.items()))
        assert key.startswith(f"procurement/{org}/") and key.endswith(".xlsx")
        assert openpyxl.load_workbook(io.BytesIO(data)).sheetnames == [
            "Summary",
            "Comparison",
        ]
        attachment = timeline.await_args.kwargs["payload"]["attachments"][0]
        assert attachment["file"] == "xlsx"

    async def test_register_updates_and_lists_only_this_workspace(
        self, flag_on, db_session, storage, converter, timeline, charges
    ):
        mine = await _org(db_session)
        theirs = await _org(db_session)
        drafted = await tools.run(
            tools.DRAFT,
            organization_id=mine,
            arguments={
                "kind": "purchase_order",
                "values": _complete_values(),
                "items": _items(),
            },
            ref_id="t:r",
        )
        refused = await tools.run(
            tools.UPDATE,
            organization_id=theirs,
            arguments={"register_id": drafted["register_id"], "status": "cancelled"},
            ref_id="t:u1",
        )
        assert refused["status"] == "error"
        listed = await tools.run(
            tools.LIST, organization_id=theirs, arguments={}, ref_id="t:l1"
        )
        assert listed["entries"] == []

        moved = await tools.run(
            tools.UPDATE,
            organization_id=mine,
            arguments={
                "number": drafted["number"],
                "status": "issued",
                "note": "Approved by MD",
            },
            ref_id="t:u2",
        )
        assert moved["entry"]["status"] == "issued"
        listed = await tools.run(
            tools.LIST,
            organization_id=mine,
            arguments={"status": "issued"},
            ref_id="t:l2",
        )
        assert [e["number"] for e in listed["entries"]] == [drafted["number"]]

    async def test_read_document_returns_text_and_tables(
        self, flag_on, storage, charges
    ):
        document = docx.Document()
        document.add_paragraph("Quotation Q-2291")
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Item", "Rate"
        table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Cement", "380"
        out = io.BytesIO()
        document.save(out)
        storage.files["knowledge_base/7/u/quote.docx"] = out.getvalue()
        row = SimpleNamespace(
            document_uuid="11111111-2222-3333-4444-555555555555",
            filename="quote.docx",
            mime_type="",
            custom_metadata={"s3_key": "knowledge_base/7/u/quote.docx"},
        )
        from api.db import db_client

        with patch.object(
            db_client, "get_document_by_uuid", new=AsyncMock(return_value=row)
        ) as lookup:
            result = await tools.run(
                tools.READ,
                organization_id=7,
                arguments={"document": row.document_uuid},
                ref_id="t:rd",
            )
        assert lookup.await_args.args == (row.document_uuid, 7)
        assert result["status"] == "success"
        assert "Quotation Q-2291" in result["text"]
        assert result["tables"] == [[["Item", "Rate"], ["Cement", "380"]]]

    async def test_gmail_not_connected_offers_the_connect_card(self, flag_on):
        with (
            patch(
                "api.services.integrations.composio.client.connected_accounts",
                new=AsyncMock(return_value=[]),
            ),
            patch(
                "api.services.workflow.connector_offer.offer",
                new=AsyncMock(return_value={"status": "offered"}),
            ) as offer,
        ):
            result = await tools.run(
                tools.SAVE_ATTACHMENT,
                organization_id=7,
                arguments={"message_id": "m1", "filename": "quote.pdf"},
                ref_id="t:g",
            )
        assert result["status"] == "needs_connection"
        assert result["app"] == "gmail"
        assert offer.await_args.kwargs["arguments"]["app"] == "gmail"

    async def test_a_gmail_attachment_is_saved_as_an_upload(self, flag_on, charges):
        accounts = [{"app": "gmail", "connected_account_id": "ca_1"}]
        calls = []

        async def execute(**kwargs):
            calls.append(kwargs)
            if kwargs["tool_slug"] == "GMAIL_FETCH_MESSAGE_BY_MESSAGE_ID":
                return {
                    "status": "success",
                    "data": {
                        "attachmentList": [
                            {"attachmentId": "att-9", "filename": "Quote.pdf"}
                        ]
                    },
                }
            return {
                "status": "success",
                "data": {
                    "file": {
                        "name": "Quote.pdf",
                        "mimetype": "application/pdf",
                        "data": "JVBERi0xLjQKJXRlc3QK",
                    }
                },
            }

        with (
            patch(
                "api.services.integrations.composio.client.connected_accounts",
                new=AsyncMock(return_value=accounts),
            ),
            patch(
                "api.services.integrations.composio.client.execute_tool", new=execute
            ),
            patch(
                "api.services.documents.sources.store_as_upload",
                new=AsyncMock(
                    return_value={
                        "document_uuid": "u-1",
                        "filename": "Quote.pdf",
                        "size_bytes": 8,
                    }
                ),
            ) as stored,
        ):
            result = await tools.run(
                tools.SAVE_ATTACHMENT,
                organization_id=7,
                arguments={"message_id": "m1", "filename": "quote.pdf"},
                ref_id="t:g2",
            )
        assert result["status"] == "success", result
        assert calls[-1]["tool_slug"] == "GMAIL_GET_ATTACHMENT"
        assert calls[-1]["arguments"]["attachment_id"] == "att-9"
        assert calls[-1]["connected_account_id"] == "ca_1"
        assert stored.await_args.kwargs["data"].startswith(b"%PDF")


@pytest.mark.asyncio
class TestFlagOff:
    async def test_the_tools_refuse(self, monkeypatch):
        monkeypatch.setattr(constants, "PROCUREMENT_DOCS_2026_09_ENABLED", False)
        result = await tools.run(
            tools.LIST_FIELDS,
            organization_id=1,
            arguments={"template": "rfq"},
            ref_id="x",
        )
        assert result["status"] == "unavailable"
