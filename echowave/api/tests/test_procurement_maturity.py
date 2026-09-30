"""The procurement maturity assessor (KAN-182, E-2).

What these hold: the seven dimensions are fixed and the tool does the
arithmetic; a score without a reason, or outside 1 to 5, is refused by
name; the band follows the mean; gaps and the roadmap follow the scores
with the lowest first; a high score with no document read is flagged; the
report and the score sheet say what the numbers say; the tool files one
register row and hands over three files, and a bad answer produces nothing.
"""

from __future__ import annotations

import io
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import docx
import pytest
from openpyxl import load_workbook

from api.db import db_client as client
from api.services.documents import maturity, spreadsheet, tools

REASON = "Seen in the interview and in the two purchase orders read."


def _answers(**scores) -> dict:
    base = {k: 3 for k in maturity.KEYS}
    base.update(scores)
    return {
        k: {"score": v, "reason": REASON, "evidence": "PO-0012.pdf"}
        for k, v in base.items()
    }


class TestTheDimensions:
    def test_there_are_seven_each_with_agents_that_close_it(self):
        assert len(maturity.DIMENSIONS) == 7
        assert maturity.KEYS == (
            "organisation",
            "policy",
            "sourcing",
            "systems",
            "suppliers",
            "performance",
            "talent",
        )
        assert all(d.agents for d in maturity.DIMENSIONS)


class TestScoring:
    def test_the_overall_is_the_mean_and_the_band_follows_it(self):
        out = maturity.assess("Shreeram", _answers(organisation=5, policy=4))
        assert out.overall == 3.4
        assert out.band == "Defined"
        assert maturity.band_of(4.5)[0] == "Optimised"
        assert maturity.band_of(1.2)[0] == "Ad hoc"

    def test_gaps_partial_and_the_roadmap_lowest_first(self):
        out = maturity.assess(
            "Shreeram", _answers(policy=1, systems=2, talent=5, sourcing=4)
        )
        assert out.gaps == ("Policy and governance", "Tools and systems")
        assert "Organisation" in out.partial and "Talent and skills" not in out.partial
        waves = [(w.days, w.dimension) for w in out.roadmap]
        assert waves[0] == ("Days 1-30", "Policy and governance")
        assert waves[1] == ("Days 31-60", "Tools and systems")
        assert all(d == "Days 61-90" for d, _ in waves[2:])
        assert "Talent and skills" not in [d for _, d in waves]
        assert out.roadmap[0].agents == ("Approval matrix", "Audit log")

    def test_a_high_score_with_nothing_read_is_flagged(self):
        answers = _answers(suppliers=4)
        answers["suppliers"]["evidence"] = ""
        out = maturity.assess("Shreeram", answers)
        assert any("Supplier management scored 4" in w for w in out.warnings)
        assert not [w for w in out.warnings if "Organisation" in w]

    @pytest.mark.parametrize(
        "answers, says",
        [
            ({}, "missing"),
            (
                {**_answers(), "budget": {"score": 3, "reason": REASON}},
                "Unknown dimension",
            ),
            ({**_answers(), "policy": {"score": 6, "reason": REASON}}, "from 1 to 5"),
            (
                {**_answers(), "policy": {"score": "high", "reason": REASON}},
                "whole number",
            ),
            ({**_answers(), "policy": {"score": 3, "reason": "ok"}}, "reason must say"),
            ({**_answers(), "policy": 3}, "object with score"),
        ],
    )
    def test_what_cannot_be_scored_is_refused_by_name(self, answers, says):
        with pytest.raises(maturity.AssessmentError) as exc:
            maturity.assess("Shreeram", answers)
        assert says in str(exc.value)

    def test_the_client_is_needed(self):
        with pytest.raises(maturity.AssessmentError):
            maturity.assess("  ", _answers())


class TestTheFiles:
    def test_the_report_says_what_the_numbers_say(self):
        out = maturity.assess("Shreeram Constructions", _answers(policy=1))
        data = maturity.report_docx(out, assessed_by="Nithish", on="2026-09-29")
        document = docx.Document(io.BytesIO(data))
        text = "\n".join(p.text for p in document.paragraphs)
        cells = [c.text for t in document.tables for r in t.rows for c in r.cells]
        assert "Shreeram Constructions" in text
        assert f"Overall {out.overall} of 5, {out.band}." in text
        assert "Prepared by Nithish" in text and "2026-09-29" in text
        assert "Gaps (scored 1 or 2): Policy and governance." in text
        assert "Policy and governance" in cells and "Days 1-30" in cells
        assert "Approval matrix, Audit log" in cells

    def test_the_score_sheet_has_the_scores_and_the_roadmap(self):
        out = maturity.assess("Shreeram", _answers(policy=1))
        book = load_workbook(
            io.BytesIO(spreadsheet.build_workbook("x", maturity.score_sheets(out)))
        )
        scores = book["Scores"]
        rows = list(scores.iter_rows(values_only=True))
        assert rows[0][:2] == ("Dimension", "Score (1-5)")
        assert rows[1][:2] == ("Organisation", 3)
        assert rows[-1][:3] == ("Overall", out.overall, out.band)
        roadmap = list(book["Roadmap"].iter_rows(values_only=True))
        assert roadmap[1][:2] == ("Days 1-30", "Policy and governance")


class TestTheTool:
    def test_it_is_a_document_tool_with_a_schema_and_a_rule(self):
        assert tools.ASSESS in tools.NAMES
        schema = next(s for s in tools.schemas() if s["name"] == tools.ASSESS)
        assert set(schema["parameters"]["required"]) == {"client", "answers"}
        assert "organisation" in schema["description"]
        assert tools.ASSESS in tools.RULES

    @pytest.mark.asyncio
    async def test_a_bad_answer_produces_nothing(self):
        with patch("api.services.documents.register.create", new=AsyncMock()) as create:
            out = await tools.assess_maturity(7, {"client": "X", "answers": {}})
        assert out["status"] == "invalid" and "missing" in out["error"]
        create.assert_not_awaited()

    @pytest.mark.asyncio
    async def test_it_files_one_row_and_hands_over_three_files(self):
        row = SimpleNamespace(id=31, number="DOC/26-27/0004", status="draft", data={})
        put: dict[str, bytes] = {}

        async def _put(key, data):
            put[key] = data

        class _Session:
            async def __aenter__(self):
                return self

            async def __aexit__(self, *a):
                return False

            async def commit(self):
                pass

        with (
            patch.object(client, "async_session", new=lambda: _Session()),
            patch(
                "api.services.documents.register.create",
                new=AsyncMock(return_value=row),
            ),
            patch(
                "api.services.documents.register.get", new=AsyncMock(return_value=row)
            ),
            patch("api.services.documents.tools._put", new=_put),
            patch(
                "api.services.documents.tools._link",
                new=AsyncMock(return_value="https://x/f"),
            ),
            patch(
                "api.services.documents.convert.to_pdf",
                new=AsyncMock(return_value=b"%PDF-1.4"),
            ),
            patch("api.services.documents.tools._hand_over", new=AsyncMock()) as handed,
        ):
            out = await tools.assess_maturity(
                7,
                {
                    "client": "Shreeram",
                    "answers": _answers(policy=1),
                    "assessed_by": "Nithish",
                },
                workflow_id=12,
            )
        assert out["status"] == "success"
        assert (out["number"], out["band"], out["gaps"]) == (
            "DOC/26-27/0004",
            "Defined",
            ["Policy and governance"],
        )
        assert [f["filename"] for f in out["files"]] == [
            "DOC-26-27-0004.pdf",
            "DOC-26-27-0004.docx",
            "DOC-26-27-0004 scores.xlsx",
        ]
        assert len(put) == 3 and row.status == "issued"
        assert row.docx_key and row.pdf_key and row.xlsx_key
        kw = handed.await_args.kwargs
        assert kw["workflow_id"] == 12 and len(kw["attachments"]) == 3
        assert "Shreeram" in kw["summary"] and "Defined" in kw["summary"]

    @pytest.mark.asyncio
    async def test_run_routes_to_it_and_off_it_is_unavailable(self, monkeypatch):
        from api import constants

        monkeypatch.setattr(constants, "PROCUREMENT_DOCS_2026_09_ENABLED", False)
        assert (
            await tools.run(tools.ASSESS, organization_id=7, arguments={}, ref_id="r")
        )["status"] == "unavailable"
        monkeypatch.setattr(constants, "PROCUREMENT_DOCS_2026_09_ENABLED", True)
        with patch(
            "api.services.documents.tools.assess_maturity",
            new=AsyncMock(return_value={"status": "success"}),
        ) as fn:
            with patch("api.services.documents.tools._charge", new=AsyncMock()):
                out = await tools.run(
                    tools.ASSESS,
                    organization_id=7,
                    arguments={"client": "X"},
                    ref_id="r",
                    workflow_id=3,
                )
        assert out["status"] == "success"
        assert (
            fn.await_args.args == (7, {"client": "X"})
            and fn.await_args.kwargs["workflow_id"] == 3
        )
