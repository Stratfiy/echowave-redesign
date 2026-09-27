"""An Excel workbook is a document template too (KAN-159, slice 2).

Same convention as Word: ``{{field}}`` in a cell is a field, ``{{items.col}}``
marks an item row. What differs is what a workbook carries that a Word file
does not -- formulas -- and the rule that follows: **a formula cell is never
written**, so the sheet's own arithmetic keeps working after the fill. Item
rows are the pre-made rows of the template, filled downwards and blanked
when unused, never inserted, because inserting rows would silently break
every SUM below them.

Modelled on the founder's NL/001 export invoice: six item rows with
``=ROUND(D*E,2)`` per line, ``=SUM`` under them, tax and total from those.
"""

from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
import pytest

from api.services.documents import templates


def _workbook() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws["B2"] = "NAUTOMATION LABS PRIVATE LIMITED"
    ws["E2"] = "INVOICE"
    ws["B9"] = "{{ bill_to_name }}"
    ws["C9"] = "=B9"
    ws["E9"] = "{{document_number}}"
    ws["F9"] = "{{date}}"
    ws["B17"], ws["D17"], ws["E17"], ws["F17"] = (
        "DESCRIPTION",
        "QTY",
        "UNIT PRICE",
        "AMOUNT",
    )
    for row in (18, 19, 20):
        ws[f"B{row}"] = "{{items.description}}"
        ws[f"D{row}"] = "{{items.quantity}}"
        ws[f"E{row}"] = "{{items.unit_price}}"
        ws[f"F{row}"] = f'=IF(D{row}="",ROUND(1*E{row},2),ROUND(D{row}*E{row},2))'
        ws[f"F{row}"].number_format = "#,##0.00"
    ws["F21"] = "=SUM(F18:F20)"
    ws["E23"] = "Total Amount in Words"
    ws["F23"] = "{{total_in_words}}"
    ws["B24"] = "{{lut_note}}"
    ws["B2"].font = openpyxl.styles.Font(bold=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _sheet(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data)).active


class TestInspecting:
    def test_fields_and_item_columns_in_reading_order(self):
        found = templates.inspect(_workbook())
        assert found["fields"] == [
            "bill_to_name",
            "document_number",
            "date",
            "total_in_words",
            "lut_note",
        ]
        assert found["item_columns"] == ["description", "quantity", "unit_price"]
        assert found["item_rows"] == 3

    def test_a_word_file_still_inspects_as_before(self):
        import docx

        document = docx.Document()
        document.add_paragraph("PO {{document_number}} for {{vendor_name}}")
        out = io.BytesIO()
        document.save(out)
        found = templates.inspect(out.getvalue())
        assert found["fields"] == ["document_number", "vendor_name"]
        assert "item_rows" not in found


class TestFilling:
    def test_values_land_numbers_stay_numbers_and_formulas_are_untouched(self):
        filled = templates.fill(
            _workbook(),
            {
                "bill_to_name": "Copy Hero LLC",
                "document_number": "NL/003",
                "date": "27-Sep-2026",
                "total_in_words": "Three Hundred and Sixteen Only",
                "lut_note": "SUPPLY MEANT FOR EXPORT UNDER LUT",
            },
            [
                {
                    "description": "Service fee",
                    "quantity": 1,
                    "unit_price": Decimal("250.00"),
                },
                {
                    "description": "HOSTINGER - 1 MONTH",
                    "quantity": 3,
                    "unit_price": "22.00",
                },
            ],
        )
        ws = _sheet(filled)
        assert ws["B9"].value == "Copy Hero LLC"
        assert ws["C9"].value == "=B9"
        assert ws["E9"].value == "NL/003"
        assert ws["B18"].value == "Service fee"
        assert ws["D18"].value == 1
        assert ws["E18"].value == 250.0
        assert ws["E19"].value == 22.0
        assert ws["F18"].value == '=IF(D18="",ROUND(1*E18,2),ROUND(D18*E18,2))'
        assert ws["F21"].value == "=SUM(F18:F20)"
        assert ws["F18"].number_format == "#,##0.00"
        assert ws["B2"].font.bold is True

    def test_unused_item_rows_are_blanked_not_left_as_placeholders(self):
        filled = templates.fill(
            _workbook(), {"bill_to_name": "X"}, [{"description": "Only one"}]
        )
        ws = _sheet(filled)
        assert ws["B19"].value is None
        assert ws["D20"].value is None
        # The row's own formula still stands, so a later hand-typed line adds up.
        assert ws["F20"].value.startswith("=IF(")

    def test_more_items_than_rows_is_refused_by_the_numbers(self):
        with pytest.raises(templates.TemplateError) as exc:
            templates.fill(
                _workbook(),
                {},
                [{"description": f"line {n}"} for n in range(4)],
            )
        assert "4" in str(exc.value) and "3" in str(exc.value)

    def test_an_unknown_field_stays_visible(self):
        ws = _sheet(templates.fill(_workbook(), {}, []))
        assert ws["E9"].value == "{{document_number}}"

    def test_a_template_that_is_neither_word_nor_excel_is_refused(self):
        with pytest.raises(templates.TemplateError):
            templates.inspect(b"PK\x03\x04 not really")


class TestPdf:
    @pytest.mark.asyncio
    async def test_the_fallback_renders_the_sheet_as_a_table(self, monkeypatch):
        from api import constants
        from api.services.documents import convert

        monkeypatch.setattr(constants, "GOTENBERG_URL", "")
        pdf = await convert.to_pdf(_workbook(), filename="NL-003.xlsx")
        assert pdf[:4] == b"%PDF"
