"""Word to PDF, and a PDF even when the converter is down.

A PO is sent as a PDF. Gotenberg (LibreOffice behind HTTP) makes one that
looks like the .docx; when it cannot be reached, a plain rendering of the
same text and tables is produced instead, so the draft is never left without
its PDF.
"""

import io

import docx
import httpx
import pypdf
import pytest

from api.services.documents import convert


def _docx() -> bytes:
    document = docx.Document()
    document.add_heading("Purchase Order", 0)
    document.add_paragraph("Vendor: Bharat Steels Pvt Ltd")
    table = document.add_table(rows=2, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Sl", "Description", "Amount (₹)"]):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ["1", "TMT bars 12mm", "1,18,000.00"]):
        cell.text = text
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _text(pdf: bytes) -> str:
    reader = pypdf.PdfReader(io.BytesIO(pdf))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


@pytest.mark.asyncio
class TestGotenberg:
    async def test_the_docx_goes_to_the_libreoffice_route(self):
        seen = {}

        def handler(request: httpx.Request) -> httpx.Response:
            seen["url"] = str(request.url)
            seen["body"] = request.content
            return httpx.Response(200, content=b"%PDF-1.7 converted")

        pdf = await convert.docx_to_pdf(
            _docx(),
            filename="PO-26-27-0001.docx",
            base_url="http://gotenberg:3000",
            transport=httpx.MockTransport(handler),
        )
        assert pdf == b"%PDF-1.7 converted"
        assert seen["url"] == "http://gotenberg:3000/forms/libreoffice/convert"
        assert b'filename="PO-26-27-0001.docx"' in seen["body"]

    async def test_unreachable_falls_back_to_a_plain_pdf(self, caplog):
        def handler(request: httpx.Request) -> httpx.Response:
            raise httpx.ConnectError("no route to host", request=request)

        pdf = await convert.docx_to_pdf(_docx(), transport=httpx.MockTransport(handler))
        assert pdf.startswith(b"%PDF")
        text = _text(pdf)
        assert "Purchase Order" in text
        assert "Bharat Steels" in text
        assert "TMT bars 12mm" in text
        assert "1,18,000.00" in text

    async def test_an_error_answer_falls_back_too(self):
        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(503, content=b"busy")

        pdf = await convert.docx_to_pdf(_docx(), transport=httpx.MockTransport(handler))
        assert pdf.startswith(b"%PDF")
        assert "Bharat Steels" in _text(pdf)

    async def test_a_non_pdf_answer_is_not_trusted(self):
        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(200, content=b"<html>proxy login</html>")

        pdf = await convert.docx_to_pdf(_docx(), transport=httpx.MockTransport(handler))
        assert pdf.startswith(b"%PDF")
        assert "Bharat Steels" in _text(pdf)


def test_the_fallback_renders_directly():
    pdf = convert.render_fallback_pdf(_docx())
    assert pdf.startswith(b"%PDF")
    assert "TMT bars 12mm" in _text(pdf)
