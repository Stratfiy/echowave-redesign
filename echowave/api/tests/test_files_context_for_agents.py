"""Decibyl and every agent get context from any file in Files, and cite it.

Through the real ingestion task, the real database and the real searches.
Only the boundaries are stubbed: the object store (a local file), and the
model that looks at a picture (its reply). Files are uploaded the way the
Files page uploads them -- organisation-wide, full-document mode -- and
searched by words, which is what a workspace without embeddings has.
"""

from __future__ import annotations

import base64
import io
from pathlib import Path
from types import SimpleNamespace

import pytest

from api.db import db_client
from api.db.models import OrganizationModel, UserModel
from api.services.knowledge_base import (
    citations,
    extraction,
    folders,
    upload_keys,
    versions,
    vision,
)
from api.services.workflow import agent_timeline, decibyl, files_search
from api.services.workflow.tools import knowledge_base as kb_tool
from api.tasks import knowledge_base_processing as task_module

# A 1x1 PNG: a real picture header, so the type is read from its bytes.
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)

PRICE_BOARD_READING = (
    "DESCRIPTION: A chalkboard price list outside Lakshmi Tiffin Centre.\n"
    "TEXT:\n"
    "Masala dosa: Rs 80\n"
    "Filter coffee: Rs 30\n"
    "Open 7am to 11pm"
)


async def _organization(async_session, slug: str):
    org = OrganizationModel(provider_id=f"org-{slug}", quota_decibyl_tokens=0)
    async_session.add(org)
    await async_session.flush()
    user = UserModel(provider_id=f"user-{slug}", selected_organization_id=org.id)
    async_session.add(user)
    await async_session.flush()
    return org, user


def _workbook() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    rates = book.active
    rates.title = "Rates"
    rates.append(["Price list, October 2026"])
    rates.append(["SKU", "Item", "Price", "Unit"])
    rates.append(["SKU-101", "Basmati rice", 120, "kg"])
    rates.append(["SKU-114", "Toor dal", 165, "kg"])
    rates.append(["SKU-120", "Groundnut oil", 210, "litre"])
    notes = book.create_sheet("Notes")
    notes.append(["Note"])
    notes.append(["Prices include GST"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.fixture
def files(monkeypatch, tmp_path):
    """Upload a file to Files and let the real ingestion task read it."""
    stored: dict[str, bytes] = {}

    async def fake_download(s3_key, destination):
        Path(destination).write_bytes(stored[s3_key])
        return True

    async def no_head(s3_key):
        return None

    monkeypatch.setattr(
        task_module,
        "storage_fs",
        SimpleNamespace(adownload_file=fake_download, aget_file_metadata=no_head),
    )

    # Searched by words, as on a workspace with no embeddings set up.
    async def no_embeddings(organization_id):
        return {}

    monkeypatch.setattr(files_search, "_embeddings", no_embeddings)

    counter = {"n": 0}

    async def upload(org, user, name: str, data: bytes, *, scope="org", **where):
        counter["n"] += 1
        uuid = f"00000000-0000-4000-9000-{counter['n'] + org.id * 1000:012d}"
        key = upload_keys.build_document_key(org.id, uuid, name)
        stored[key] = data
        document = await db_client.create_document(
            organization_id=org.id,
            created_by=user.id,
            filename=name,
            file_size_bytes=0,
            file_hash="",
            mime_type="application/octet-stream",
            custom_metadata={"s3_key": key},
            document_uuid=uuid,
            retrieval_mode="full_document",
            scope=scope,
            **where,
        )
        await task_module.process_knowledge_base_document(
            ctx={},
            document_id=document.id,
            s3_key=key,
            organization_id=org.id,
            created_by_provider_id=str(user.provider_id),
            retrieval_mode="full_document",
        )
        return await db_client.get_document_by_id(document.id)

    async def reupload(org, user, document, data: bytes):
        counter["n"] += 1
        uuid = f"00000000-0000-4000-9000-{counter['n'] + org.id * 1000:012d}"
        key = upload_keys.build_document_key(org.id, uuid, document.filename)
        stored[key] = data
        await versions.begin(
            document,
            organization_id=org.id,
            s3_key=key,
            user_id=user.id,
            retrieval_mode="full_document",
        )
        await task_module.process_knowledge_base_document(
            ctx={},
            document_id=document.id,
            s3_key=key,
            organization_id=org.id,
            created_by_provider_id=str(user.provider_id),
            retrieval_mode="full_document",
        )
        return await db_client.get_document_by_id(document.id)

    return SimpleNamespace(upload=upload, reupload=reupload)


@pytest.fixture
def picture_reader(monkeypatch):
    """The model's reply to a picture, and what it was shown."""
    shown = []

    async def ask(organization_id, data, kind, filename):
        shown.append((organization_id, kind, filename, len(data)))
        return PRICE_BOARD_READING, "anthropic/test-model"

    monkeypatch.setattr(vision, "ask_model", ask)
    return shown


def _texts(result) -> str:
    return " ".join(str(c.get("text") or "") for c in result.get("chunks") or [])


@pytest.mark.asyncio
class TestAnyFileGivesContext:
    async def test_a_pictures_text_is_retrievable(
        self, db_session, async_session, files, picture_reader
    ):
        org, user = await _organization(async_session, "ctx-picture")
        document = await files.upload(org, user, "board.jpg", PNG)

        assert document.processing_status == "completed", document.processing_error
        assert document.docling_metadata["extractor"] == "vision:anthropic/test-model"
        # Read by the account's own model, about this workspace.
        assert picture_reader[0][:3] == (org.id, "image/png", "board.jpg")

        result = await files_search.search(org.id, "how much is filter coffee")
        assert result["status"] == "ok"
        assert "Filter coffee: Rs 30" in _texts(result)
        assert result["chunks"][0]["citation"].startswith("board.jpg")

        # An agent given the file reads it too.
        agent = await kb_tool.retrieve_from_knowledge_base(
            query="filter coffee price",
            organization_id=org.id,
            document_uuids=[document.document_uuid],
        )
        assert "Filter coffee: Rs 30" in _texts(agent)

    async def test_a_sheet_cell_is_retrievable_by_its_header(
        self, db_session, async_session, files
    ):
        org, user = await _organization(async_session, "ctx-sheet")
        document = await files.upload(org, user, "rates.xlsx", _workbook())
        assert document.processing_status == "completed", document.processing_error

        result = await files_search.search(org.id, "toor dal price")
        assert result["status"] == "ok"
        top = result["chunks"][0]
        # The cell, under its header, with the sheet and the title as context.
        assert "Item: Toor dal; Price: 165" in top["text"]
        assert "Sheet: Rates" in top["text"]
        assert top["sheet"] == "Rates"
        assert "rates.xlsx, sheet Rates, row" in top["citation"]

        # Every passage of the sheet carries its header names and its sheet.
        chunks = await db_client.get_chunks_for_document(document.id, org.id)
        for chunk in chunks:
            assert chunk.contextualized_text.startswith("Sheet: ")
            if chunk.chunk_metadata.get("sheet") == "Rates":
                assert "SKU: " in chunk.chunk_text and "Price: " in chunk.chunk_text

    async def test_a_moved_or_renamed_file_is_still_cited_correctly(
        self, db_session, async_session, files
    ):
        org, user = await _organization(async_session, "ctx-moved")
        document = await files.upload(org, user, "rates.xlsx", _workbook())
        before = await files_search.search(org.id, "groundnut oil price")
        assert before["chunks"][0]["citation"].startswith("rates.xlsx, sheet Rates")

        pricing = await folders.create(
            org.id, name="Pricing", parent_id=None, created_by=user.id
        )
        year = await folders.create(
            org.id, name="2026", parent_id=pricing["id"], created_by=user.id
        )
        await folders.place_document(
            org.id,
            document.document_uuid,
            filename="Wholesale rates",
            file_folder_id=year["id"],
            move=True,
        )

        after = await files_search.search(org.id, "groundnut oil price")
        top = after["chunks"][0]
        assert top["filename"] == "Wholesale rates.xlsx"
        assert top["folder"] == "Pricing/2026"
        assert top["citation"].startswith(
            "Wholesale rates.xlsx (in Pricing/2026), sheet Rates, row"
        )
        # And the thread's sources name it the same way.
        assert agent_timeline.passages_sources(after)[0]["documents"][0].startswith(
            "Wholesale rates.xlsx (in Pricing/2026)"
        )

    async def test_a_re_uploaded_file_serves_its_new_content(
        self, db_session, async_session, files
    ):
        org, user = await _organization(async_session, "ctx-again")
        document = await files.upload(
            org, user, "hours.txt", b"Delivery cut-off is four in the afternoon."
        )
        updated = await files.reupload(
            org, user, document, b"Delivery cut-off moved to six in the evening."
        )
        assert versions.current(updated) == 2

        result = await files_search.search(org.id, "delivery cut-off")
        text = _texts(result)
        assert "six in the evening" in text
        assert "four in the afternoon" not in text

        agent = await kb_tool.retrieve_from_knowledge_base(
            query="delivery cut-off",
            organization_id=org.id,
            document_uuids=[document.document_uuid],
        )
        assert "six in the evening" in _texts(agent)
        assert "four in the afternoon" not in _texts(agent)

    async def test_another_org_never_sees_any_of_it(
        self, db_session, async_session, files, picture_reader
    ):
        owner, owner_user = await _organization(async_session, "ctx-owner")
        stranger, _ = await _organization(async_session, "ctx-stranger")
        sheet = await files.upload(owner, owner_user, "rates.xlsx", _workbook())
        picture = await files.upload(owner, owner_user, "board.png", PNG)

        assert (await files_search.search(stranger.id, "toor dal price"))[
            "chunks"
        ] == []
        assert (await files_search.search(stranger.id, "filter coffee"))["chunks"] == []
        # Not even when it names the owner's files outright.
        named = await kb_tool.retrieve_from_knowledge_base(
            query="toor dal price filter coffee",
            organization_id=stranger.id,
            document_uuids=[sheet.document_uuid, picture.document_uuid],
        )
        assert named["chunks"] == []
        assert (
            await db_client.keyword_search_chunks(
                organization_id=stranger.id, terms=["toor", "coffee"]
            )
            == []
        )
        assert await db_client.readable_document_uuids(stranger.id) == []


@pytest.mark.asyncio
class TestWhoReadsWhat:
    async def test_org_files_go_to_everyone_channel_and_agent_files_only_there(
        self, db_session, async_session, files
    ):
        from api.db.models import FolderModel, WorkflowModel

        org, user = await _organization(async_session, "ctx-scope")
        channel = FolderModel(organization_id=org.id, name="Front desk")
        async_session.add(channel)
        await async_session.flush()
        bot = WorkflowModel(organization_id=org.id, user_id=user.id, name="Biller")
        async_session.add(bot)
        await async_session.flush()

        everyone = await files.upload(
            org, user, "everyone.txt", b"Office closes on Sunday."
        )
        for_channel = await files.upload(
            org,
            user,
            "channel.txt",
            b"Front desk keeps the spare key on Sunday.",
            scope="channel",
            folder_id=channel.id,
        )
        for_bot = await files.upload(
            org,
            user,
            "bot.txt",
            b"Biller sends invoices on Sunday.",
            scope="bot",
            workflow_id=bot.id,
        )

        # Decibyl: the workspace's files, not one channel's or one agent's.
        readable = set(
            await files_search.db_client.readable_document_uuids(
                org.id, exclude_scopes=files_search.NOT_FOR_DECIBYL
            )
        )
        assert everyone.document_uuid in readable
        assert not {for_channel.document_uuid, for_bot.document_uuid} & readable
        decibyl_read = await files_search.search(org.id, "what happens on sunday")
        assert "Office closes" in _texts(decibyl_read)
        assert "spare key" not in _texts(decibyl_read)
        assert "invoices" not in _texts(decibyl_read)

        # The agent: company files plus its own, never another's channel.
        agent_docs = await db_client.scoped_document_uuids(
            org.id, workflow_id=bot.id, folder_id=None
        )
        assert set(agent_docs) == {everyone.document_uuid, for_bot.document_uuid}
        # The channel's agents: company files plus the channel's.
        channel_docs = await db_client.scoped_document_uuids(
            org.id, workflow_id=None, folder_id=channel.id
        )
        assert set(channel_docs) == {everyone.document_uuid, for_channel.document_uuid}


@pytest.mark.asyncio
class TestTheThreadShowsWhichFile:
    async def test_decibyls_reply_names_the_file_folder_and_place(
        self, db_session, async_session, files, monkeypatch
    ):
        org, user = await _organization(async_session, "ctx-thread")
        await files.upload(org, user, "rates.xlsx", _workbook())
        knowledge = await files_search.search(org.id, "basmati rice price")

        block = decibyl.knowledge_block(knowledge)
        assert "(rates.xlsx, sheet Rates, row" in block
        sources = decibyl.sources_read(bots=0, facts=0, knowledge=knowledge)
        files_row = next(s for s in sources if s["kind"] == "knowledge")
        assert files_row["documents"][0].startswith("rates.xlsx, sheet Rates, row")

        recorded = []

        async def record(**kwargs):
            recorded.append(kwargs)

        monkeypatch.setattr(agent_timeline, "record_activity", record)
        answer = await files_search.for_thread(org.id, {"query": "toor dal price"})
        assert answer["status"] == "success" and answer["found"] is True
        assert "Price: 165" in answer["passages"][0]["text"]
        assert recorded[0]["payload"]["sources"][0]["documents"][0].startswith(
            "rates.xlsx, sheet Rates, row"
        )

    def test_search_files_is_offered_and_explained(self):
        names = [t["name"] for t in decibyl.office_tools()]
        assert files_search.TOOL_NAME in names
        assert files_search.TOOL_NAME in decibyl.SYSTEM


class TestReadingWell:
    def test_a_word_table_sits_under_its_own_heading(self, tmp_path):
        import docx

        document = docx.Document()
        document.add_heading("Shipping rates", level=1)
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Zone"
        table.rows[0].cells[1].text = "Rate"
        table.rows[1].cells[0].text = "South"
        table.rows[1].cells[1].text = "Rs 90"
        document.add_heading("Returns", level=1)
        document.add_paragraph("Returns are accepted for seven days.")
        path = tmp_path / "policy.docx"
        document.save(path)

        extracted = extraction.extract_document(str(path), "policy.docx")
        row = next(b for b in extracted.blocks if "Zone: South" in b.text)
        assert row.heading_path == ("Shipping rates",)
        assert row.text == "Zone: South; Rate: Rs 90"

    def test_citations_name_the_page_or_the_sheet(self):
        assert citations.location({"pages": [3]}) == "page 3"
        assert citations.location({"pages": [4, 3]}) == "pages 3-4"
        assert citations.location({"sheet": "Rates", "rows": [9, 4]}) == (
            "sheet Rates, rows 4-9"
        )
        assert citations.cite("a.pdf", "", {"pages": [2]}) == "a.pdf, page 2"
        assert citations.cite("a.pdf", "HR", {}) == "a.pdf (in HR)"

    def test_a_picture_goes_to_the_model_as_a_picture(self, monkeypatch):
        """Through the platform's own model client, as pictures on a tool
        result -- the way Studio shows a model its screenshots -- not a new
        provider."""
        import asyncio

        from api.services.agent_builder import client, settings

        seen = {}

        async def resolve(session, choice, *, organization_id):
            seen["org"] = organization_id
            return SimpleNamespace(provider="anthropic", model="m", api_key="k")

        async def complete(**kwargs):
            seen.update(kwargs)
            return client.ModelReply(text=PRICE_BOARD_READING)

        monkeypatch.setattr(settings, "resolve_for_organization", resolve)
        monkeypatch.setattr(client, "complete", complete)

        reply, label = asyncio.run(vision.ask_model(7, PNG, "image/png", "board.png"))
        assert reply == PRICE_BOARD_READING
        assert label == "anthropic/m" and seen["org"] == 7
        tool_result = seen["conversation"].messages[-1]
        assert tool_result["role"] == "tool"
        image = tool_result["content"][client.IMAGES_KEY][0]
        assert image["media_type"] == "image/png"
        assert base64.b64decode(image["data"]) == PNG

    def test_a_reading_becomes_a_description_and_lines_of_text(self):
        description, lines = vision.parse_reading(PRICE_BOARD_READING)
        assert description.startswith("A chalkboard price list")
        assert lines == [
            "Masala dosa: Rs 80",
            "Filter coffee: Rs 30",
            "Open 7am to 11pm",
        ]
        assert vision.parse_reading("DESCRIPTION: A blank wall.\nTEXT:\nNONE") == (
            "A blank wall.",
            [],
        )

    def test_a_file_named_like_a_picture_that_is_not_one_is_refused(self):
        from api.services.knowledge_base.errors import DocumentExtractionError

        with pytest.raises(DocumentExtractionError):
            vision.media_type(b"%PDF-1.7", "scan.jpg")
