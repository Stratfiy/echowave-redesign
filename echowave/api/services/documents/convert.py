"""A drafted .docx as a PDF: Gotenberg first, a plain rendering if not.

**Gotenberg** is LibreOffice behind an HTTP API, run as its own container
on the compose network (``GOTENBERG_URL``, internal only). Its LibreOffice
route turns the .docx into a PDF that looks like the Word file: same
fonts, same table borders, same page breaks.

**The fallback.** A PO is sent as a PDF, so a converter being down must not
leave a draft without one. When Gotenberg cannot be reached, answers with an
error, or answers with something that is not a PDF (a proxy's login page,
say), the document's paragraphs and tables are rendered with reportlab
instead, in body order, and a warning is logged. Plainer, never missing.
"""

from __future__ import annotations

import io
import os
from typing import Any
from xml.sax.saxutils import escape

import docx
import httpx
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from loguru import logger

from api import constants

LIBREOFFICE_ROUTE = "/forms/libreoffice/convert"
TIMEOUT_SECS = 60.0
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def to_pdf(
    data: bytes,
    *,
    filename: str = "document.docx",
    base_url: str | None = None,
    transport: httpx.AsyncBaseTransport | None = None,
) -> bytes:
    """The PDF of a .docx or .xlsx, told apart by the filename. Never raises
    for a converter problem: the fallback rendering is returned instead."""
    is_sheet = filename.lower().endswith(".xlsx")
    mime = XLSX_MIME if is_sheet else DOCX_MIME
    url = (base_url or constants.GOTENBERG_URL or "").rstrip("/")
    if url:
        try:
            async with httpx.AsyncClient(
                timeout=TIMEOUT_SECS, transport=transport
            ) as client:
                response = await client.post(
                    f"{url}{LIBREOFFICE_ROUTE}",
                    files={"files": (filename, data, mime)},
                )
            if response.status_code == 200 and response.content.startswith(b"%PDF"):
                return response.content
            logger.warning(
                "Gotenberg answered {} ({} bytes, not a PDF) for {}; rendering "
                "a plain PDF instead",
                response.status_code,
                len(response.content),
                filename,
            )
        except httpx.HTTPError as exc:
            logger.warning(
                "Gotenberg at {} could not convert {} ({}); rendering a plain "
                "PDF instead",
                url,
                filename,
                exc,
            )
    return render_fallback_sheet_pdf(data) if is_sheet else render_fallback_pdf(data)


#: The name the drafting tool has always called; a .docx is what it sends.
docx_to_pdf = to_pdf


# ---------------------------------------------------------------------------
# The fallback rendering

_FONT_CANDIDATES = (
    (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    ),
    (
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
    ),
)
_fonts: tuple[str, str, bool] | None = None


def _font_names() -> tuple[str, str, bool]:
    """(regular, bold, has the rupee sign). A TrueType font with the ₹
    glyph when the image has one; Helvetica, with ₹ written as Rs., when
    not."""
    global _fonts
    if _fonts is not None:
        return _fonts
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for regular, bold in _FONT_CANDIDATES:
        if os.path.exists(regular) and os.path.exists(bold):
            try:
                pdfmetrics.registerFont(TTFont("DocSans", regular))
                pdfmetrics.registerFont(TTFont("DocSans-Bold", bold))
                _fonts = ("DocSans", "DocSans-Bold", "dejavu" in regular.lower())
                return _fonts
            except Exception as exc:  # noqa: BLE001 - fall to Helvetica
                logger.warning("Could not register {}: {}", regular, exc)
    _fonts = ("Helvetica", "Helvetica-Bold", False)
    return _fonts


def _body_blocks(document) -> list[Any]:
    blocks: list[Any] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            blocks.append(DocxParagraph(child, document))
        elif child.tag == qn("w:tbl"):
            blocks.append(DocxTable(child, document))
    return blocks


def _cell_fill(cell) -> str | None:
    """A cell's shading as a hex colour, when it has one."""
    shading = cell._tc.find(f"{qn('w:tcPr')}/{qn('w:shd')}")
    fill = shading.get(qn("w:fill")) if shading is not None else None
    return None if not fill or fill.lower() == "auto" else f"#{fill}"


def _borderless(table) -> bool:
    borders = table._tbl.find(f"{qn('w:tblPr')}/{qn('w:tblBorders')}")
    if borders is None:
        return False
    return all(edge.get(qn("w:val")) in ("nil", "none") for edge in borders)


def _grid_widths(table) -> list[int]:
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is None:
        return []
    widths = []
    for column in grid.findall(qn("w:gridCol")):
        try:
            widths.append(int(column.get(qn("w:w")) or 0))
        except ValueError:
            widths.append(0)
    return widths


def render_fallback_pdf(data: bytes) -> bytes:
    """Paragraphs and tables of a .docx, in order, as a plain A4 PDF: the
    first page header on top, the footer at the foot, the table widths,
    shading and bold of the Word file where it states them."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    regular, bold, rupee = _font_names()
    usable = 180 * mm

    def clean(text: str) -> str:
        text = text if rupee else text.replace("₹", "Rs.")
        return escape(text).replace("\n", "<br/>")

    base = ParagraphStyle("base", fontName=regular, fontSize=9.5, leading=12.5)
    title_style = ParagraphStyle(
        "title", parent=base, fontName=bold, fontSize=15, leading=19, spaceAfter=6
    )
    heading_style = ParagraphStyle(
        "heading", parent=base, fontName=bold, fontSize=11, leading=14, spaceBefore=4
    )
    aligns = {1: TA_CENTER, 2: TA_RIGHT}

    def paragraph_style(paragraph, parent) -> ParagraphStyle:
        runs = [r for r in paragraph.runs if r.text.strip()]
        is_bold = bool(runs) and all(r.bold for r in runs)
        sizes = [r.font.size.pt for r in runs if r.font.size is not None]
        size = max(sizes) if sizes else parent.fontSize
        alignment = paragraph.alignment
        return ParagraphStyle(
            "p",
            parent=parent,
            fontName=bold if is_bold else parent.fontName,
            fontSize=size,
            leading=size * 1.3,
            alignment=aligns.get(int(alignment), TA_LEFT)
            if alignment is not None
            else TA_LEFT,
        )

    def paragraph_flowable(paragraph, parent=base):
        text = paragraph.text
        if not text.strip():
            return None
        style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style_name == "Title":
            style = title_style
        elif style_name.startswith("Heading"):
            style = heading_style
        else:
            style = paragraph_style(paragraph, parent)
        return Paragraph(clean(text), style)

    document = docx.Document(io.BytesIO(data))
    story: list[Any] = []

    section = document.sections[0]
    for paragraph in (
        section.header.paragraphs if not section.header.is_linked_to_previous else []
    ):
        flowable = paragraph_flowable(paragraph)
        if flowable is not None:
            story.append(flowable)
    if story:
        story.append(Spacer(1, 8))

    cell_base = ParagraphStyle("cell", parent=base, fontSize=8.5, leading=10.5)
    for block in _body_blocks(document):
        if isinstance(block, DocxParagraph):
            flowable = paragraph_flowable(block)
            story.append(flowable if flowable is not None else Spacer(1, 4))
            continue

        rows: list[list[Any]] = []
        shading: list[tuple] = []
        for row_index, row in enumerate(block.rows):
            cells: list[Any] = []
            seen: set[int] = set()
            for col_index, cell in enumerate(row.cells):
                # A merged cell is returned once per grid column; print it once.
                if id(cell._tc) in seen:
                    cells.append("")
                    continue
                seen.add(id(cell._tc))
                # An empty line in a cell is room left on purpose -- the
                # space above a signature -- so it is kept as space.
                parts = [
                    paragraph_flowable(p, cell_base) or Spacer(1, 9)
                    for p in cell.paragraphs
                ]
                cells.append(parts or "")
                fill = _cell_fill(cell)
                if fill:
                    shading.append(
                        (
                            "BACKGROUND",
                            (col_index, row_index),
                            (col_index, row_index),
                            colors.HexColor(fill),
                        )
                    )
            rows.append(cells)
        if not rows:
            continue
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        grid = _grid_widths(block)[:width]
        if len(grid) == width and sum(grid) > 0:
            total = sum(grid)
            scale = min(1.0, (usable / mm) / (total / 56.7))
            col_widths = [w / 56.7 * mm * scale for w in grid]
        else:
            col_widths = [usable / width] * width
        commands: list[tuple] = [
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            *shading,
        ]
        if not _borderless(block):
            commands.append(("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa1a9")))
        alignment = block.alignment
        h_align = (
            {1: "CENTER", 2: "RIGHT"}.get(int(alignment), "LEFT")
            if alignment is not None
            else "LEFT"
        )
        table = Table(
            rows, repeatRows=1 if shading else 0, hAlign=h_align, colWidths=col_widths
        )
        table.setStyle(TableStyle(commands))
        story.append(table)
        story.append(Spacer(1, 6))

    footer_lines = (
        [p.text for p in section.footer.paragraphs if p.text.strip()]
        if not section.footer.is_linked_to_previous
        else []
    )
    footer_text = clean(" ".join(footer_lines).replace("Page", "").strip(" ·"))
    footer_style = ParagraphStyle(
        "footer",
        parent=base,
        fontSize=7.5,
        textColor=colors.HexColor("#5f666d"),
        alignment=TA_CENTER,
    )

    def draw_footer(canvas, doc) -> None:
        if not footer_text:
            return
        canvas.saveState()
        flowable = Paragraph(f"{footer_text} · Page {doc.page}", footer_style)
        _, height = flowable.wrap(usable, 20 * mm)
        flowable.drawOn(canvas, 15 * mm, 8 * mm)
        canvas.restoreState()

    if not story:
        story.append(Paragraph(" ", base))
    out = io.BytesIO()
    SimpleDocTemplate(
        out,
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=18 * mm,
        title="Document",
    ).build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
    return out.getvalue()


__all__ = ["LIBREOFFICE_ROUTE", "docx_to_pdf", "render_fallback_pdf"]


def render_fallback_sheet_pdf(data: bytes) -> bytes:
    """Each sheet of a workbook as a plain table, one after the other. The
    workbook is never calculated here, so a formula cell shows its formula;
    the converter, when it is up, shows the values."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    from api.services.documents import xlsx_templates

    regular, bold, _ = _font_names()
    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    cell_style = ParagraphStyle("cell", fontName=regular, fontSize=8, leading=10)
    title_style = ParagraphStyle("title", fontName=bold, fontSize=11, leading=14)
    story: list[Any] = []
    for title, rows in xlsx_templates.rows_as_text(data):
        story.append(Paragraph(escape(title), title_style))
        story.append(Spacer(1, 3 * mm))
        width = max(len(r) for r in rows)
        body = [
            [Paragraph(escape(c), cell_style) for c in r + [""] * (width - len(r))]
            for r in rows
        ]
        table = Table(body, repeatRows=0)
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(table)
        story.append(Spacer(1, 6 * mm))
    if not story:
        story.append(Paragraph("(empty workbook)", cell_style))
    doc.build(story)
    return out.getvalue()
