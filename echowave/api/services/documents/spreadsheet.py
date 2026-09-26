"""Spreadsheets an agent hands over: any table, and the cost-bid analysis.

**What every sheet gets.** A bold, shaded heading row frozen in place; rupee
formats with Indian grouping (₹12,34,567.00) on money columns; widths that
fit the longest value; and any formulas asked for (a ``SUM`` under a
column, a ``MIN`` across bids) written as formulas, so the sheet keeps
working when somebody changes a rate.

**Money columns** are the ones declared ``{"name": …, "type": "money"}``,
or whose heading says so (₹, amount, rate, total, value, price, cost).
``number_formats`` overrides either.

**Cost-bid analysis.** Rates per item per vendor become landed cost per unit
-- rate plus freight, plus GST on both -- and line totals, per-vendor totals
and an L1/L2/L3 ranking. The lowest landed cost on each line is shaded. A
vendor who did not quote every line is ranked after every vendor who did,
and says so: the cheapest incomplete bid is not the cheapest bid.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any

from api.services.documents import money

#: ₹ with Indian grouping: crores, lakhs, then thousands.
INR_FORMAT = (
    '[>=10000000]"₹"#\\,##\\,##\\,##0.00;[>=100000]"₹"#\\,##\\,##0.00;"₹"#,##0.00'
)
HEADER_FILL = "EEF0F2"
LOWEST_FILL = "D9EAD3"
MAX_ROWS = 5_000
MAX_SHEETS = 10
_MONEY_WORDS = re.compile(
    r"(₹|\bamount\b|\brate\b|\btotal\b|\bvalue\b|\bprice\b|\bcost\b|\binr\b)",
    re.IGNORECASE,
)
_CELL = re.compile(r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")
_BAD_SHEET_CHARS = re.compile(r"[\\/?*\[\]:]")


class SpreadsheetError(ValueError):
    """Said to the model as it is."""


def _column_spec(column: Any) -> tuple[str, str | None]:
    if isinstance(column, dict):
        return str(column.get("name") or ""), (str(column.get("type") or "") or None)
    return str(column), None


def _sheet_name(name: Any, index: int, used: set[str]) -> str:
    text = _BAD_SHEET_CHARS.sub("-", str(name or "")).strip()[:31] or f"Sheet{index}"
    base, n = text, 2
    while text in used:
        suffix = f" ({n})"
        text = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(text)
    return text


def _cell_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (int, float, str)) or value is None:
        return value
    return str(value)


def build_workbook(title: str, sheets: list[dict[str, Any]]) -> bytes:
    """An .xlsx from sheet specs. See the module docstring for the rules."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if not sheets:
        raise SpreadsheetError("Give at least one sheet.")
    if len(sheets) > MAX_SHEETS:
        raise SpreadsheetError(f"At most {MAX_SHEETS} sheets.")
    book = Workbook()
    book.remove(book.active)
    book.properties.title = str(title or "Spreadsheet")[:200]
    used: set[str] = set()
    for index, spec in enumerate(sheets, start=1):
        if not isinstance(spec, dict):
            raise SpreadsheetError(
                "Each sheet is an object with name, columns and rows."
            )
        sheet = book.create_sheet(_sheet_name(spec.get("name"), index, used))
        columns = [_column_spec(c) for c in list(spec.get("columns") or [])]
        rows = list(spec.get("rows") or [])
        if len(rows) > MAX_ROWS:
            raise SpreadsheetError(f"At most {MAX_ROWS} rows a sheet.")
        formats = {
            str(k): str(v) for k, v in dict(spec.get("number_formats") or {}).items()
        }
        money_columns = {str(c) for c in list(spec.get("money_columns") or [])}

        column_formats: list[str | None] = []
        for name, kind in columns:
            if name in formats:
                column_formats.append(formats[name])
            elif (
                kind == "money"
                or name in money_columns
                or (kind is None and _MONEY_WORDS.search(name))
            ):
                column_formats.append(INR_FORMAT)
            elif kind == "percent":
                column_formats.append("0.00%")
            else:
                column_formats.append(None)

        header_font = Font(bold=True)
        header_fill = PatternFill("solid", fgColor=HEADER_FILL)
        for col, (name, _) in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=col, value=name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        widths = [len(name) for name, _ in columns]

        for r, row in enumerate(rows, start=2):
            values = list(row) if isinstance(row, (list, tuple)) else [row]
            for c, value in enumerate(values, start=1):
                cell = sheet.cell(row=r, column=c, value=_cell_value(value))
                fmt = column_formats[c - 1] if c - 1 < len(column_formats) else None
                if fmt and isinstance(cell.value, (int, float)):
                    cell.number_format = fmt
                shown = (
                    money.format_inr(value)
                    if fmt == INR_FORMAT and isinstance(cell.value, (int, float))
                    else str(value if value is not None else "")
                )
                while len(widths) < c:
                    widths.append(0)
                widths[c - 1] = max(
                    widths[c - 1], len(shown) + (2 if fmt == INR_FORMAT else 0)
                )

        for formula in list(spec.get("formulas") or []):
            if not isinstance(formula, dict):
                raise SpreadsheetError("A formula is {cell, formula}.")
            ref = str(formula.get("cell") or "").strip().upper()
            text = str(formula.get("formula") or "").strip()
            if not _CELL.match(ref) or not text:
                raise SpreadsheetError(
                    f"{formula.get('cell')!r} is not a cell like D12."
                )
            cell = sheet[ref]
            cell.value = text if text.startswith("=") else f"={text}"
            fmt = formula.get("number_format")
            if fmt:
                cell.number_format = str(fmt)
            elif (
                cell.column - 1 < len(column_formats)
                and column_formats[cell.column - 1]
            ):
                cell.number_format = column_formats[cell.column - 1]
            if formula.get("bold"):
                cell.font = Font(bold=True)

        for ref in list(spec.get("highlights") or []):
            sheet[str(ref)].fill = PatternFill("solid", fgColor=LOWEST_FILL)
        for row_number in list(spec.get("bold_rows") or []):
            for cell in sheet[int(row_number)]:
                cell.font = Font(bold=True)

        for col, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = max(
                8, min(60, width + 2)
            )
        for key, width in dict(spec.get("column_widths") or {}).items():
            letter = (
                str(key).upper() if re.fullmatch(r"[A-Za-z]{1,3}", str(key)) else None
            )
            if letter is None:
                names = [n for n, _ in columns]
                if str(key) in names:
                    letter = get_column_letter(names.index(str(key)) + 1)
            if letter:
                try:
                    sheet.column_dimensions[letter].width = max(
                        4, min(120, float(width))
                    )
                except (TypeError, ValueError):
                    raise SpreadsheetError("column_widths are numbers.") from None
        sheet.freeze_panes = "A2"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Cost-bid analysis


@dataclass
class BidAnalysis:
    sheets: list[dict[str, Any]]
    ranking: list[dict[str, Any]] = field(default_factory=list)
    #: The vendor with the lowest landed cost on each item, by item index.
    lowest: list[str | None] = field(default_factory=list)


def _per_item(value: Any, count: int, name: str, what: str) -> list[Decimal]:
    if isinstance(value, (list, tuple)):
        if len(value) != count:
            raise SpreadsheetError(
                f"{name}: {what} needs one value per item ({count})."
            )
        return [money.dec(v or 0) for v in value]
    return [money.dec(value or 0)] * count


def _letter(col: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(col)


def cost_bid_analysis(
    vendors: list[dict[str, Any]], items: list[dict[str, Any]]
) -> BidAnalysis:
    """Landed cost per line per vendor, totals, L1/L2/L3 and a summary.

    ``items``: ``[{description, qty, unit}]``. ``vendors``: ``[{name,
    rates: [one per item, null where not quoted], freight?: per unit, one
    number or one per item, gst_rate?: %, one or one per item, gstin?}]``.
    """
    if not items:
        raise SpreadsheetError("Give the items being compared.")
    if not vendors:
        raise SpreadsheetError("Give at least one vendor's rates.")
    count = len(items)
    quantities = []
    for index, item in enumerate(items, start=1):
        try:
            qty = money.dec(item.get("qty") if isinstance(item, dict) else None)
        except money.MoneyError:
            raise SpreadsheetError(f"Item {index} needs a quantity.") from None
        if qty < 0:
            raise SpreadsheetError(f"Item {index}: quantity cannot be negative.")
        quantities.append(qty)

    table: list[dict[str, Any]] = []
    for vendor in vendors:
        name = str(vendor.get("name") or "").strip()
        if not name:
            raise SpreadsheetError("Every vendor needs a name.")
        rates_raw = vendor.get("rates")
        if not isinstance(rates_raw, (list, tuple)) or len(rates_raw) != count:
            raise SpreadsheetError(
                f"{name}: give one rate per item ({count}), null if not quoted."
            )
        freight = _per_item(vendor.get("freight"), count, name, "freight")
        gst = _per_item(vendor.get("gst_rate", 18), count, name, "gst_rate")
        rates: list[Decimal | None] = []
        landed: list[Decimal | None] = []
        totals: list[Decimal | None] = []
        for i, raw in enumerate(rates_raw):
            if raw is None or (isinstance(raw, str) and not raw.strip()):
                rates.append(None)
                landed.append(None)
                totals.append(None)
                continue
            rate = money.dec(raw)
            if rate < 0 or freight[i] < 0 or gst[i] < 0:
                raise SpreadsheetError(
                    f"{name}: rates, freight and GST cannot be negative."
                )
            unit = money.to_paise((rate + freight[i]) * (100 + gst[i]) / 100)
            rates.append(rate)
            landed.append(unit)
            totals.append(money.to_paise(unit * quantities[i]))
        quoted = sum(1 for t in totals if t is not None)
        table.append(
            {
                "name": name,
                "gstin": str(vendor.get("gstin") or ""),
                "rates": rates,
                "landed": landed,
                "totals": totals,
                "quoted": quoted,
                "complete": quoted == count,
                "total": sum((t for t in totals if t is not None), money.ZERO),
            }
        )

    lowest: list[str | None] = []
    for i in range(count):
        options = [
            (v["landed"][i], v["name"]) for v in table if v["landed"][i] is not None
        ]
        lowest.append(min(options)[1] if options else None)

    ordered = sorted(table, key=lambda v: (not v["complete"], v["total"]))
    best = ordered[0]["total"] if ordered and ordered[0]["complete"] else None
    ranking = []
    for position, vendor in enumerate(ordered, start=1):
        ranking.append(
            {
                "rank": f"L{position}",
                "vendor": vendor["name"],
                "gstin": vendor["gstin"],
                "total": vendor["total"],
                "items_quoted": f"{vendor['quoted']} of {count}",
                "complete": vendor["complete"],
            }
        )

    # The comparison sheet: four item columns, three per vendor, one MIN.
    columns: list[Any] = ["Sl", "Item", "Qty", "Unit"]
    for vendor in table:
        columns += [
            {"name": f"{vendor['name']} rate (₹)", "type": "money"},
            {"name": f"{vendor['name']} landed/unit (₹)", "type": "money"},
            {"name": f"{vendor['name']} total (₹)", "type": "money"},
        ]
    columns += [{"name": "Lowest landed/unit (₹)", "type": "money"}, "L1 on this item"]
    rows = []
    highlights = []
    formulas = []
    first_row, last_row = 2, count + 1
    for i, item in enumerate(items):
        row: list[Any] = [
            i + 1,
            str(item.get("description") or f"Item {i + 1}"),
            float(quantities[i]),
            str(item.get("unit") or ""),
        ]
        landed_cells = []
        for v_index, vendor in enumerate(table):
            base_col = 5 + v_index * 3
            row += [vendor["rates"][i], vendor["landed"][i], vendor["totals"][i]]
            landed_cells.append(f"{_letter(base_col + 1)}{i + 2}")
            if lowest[i] == vendor["name"]:
                highlights += [f"{_letter(base_col + c)}{i + 2}" for c in range(3)]
        min_col = 5 + len(table) * 3
        row += [None, lowest[i] or "No quote"]
        formulas.append(
            {
                "cell": f"{_letter(min_col)}{i + 2}",
                "formula": f"MIN({','.join(landed_cells)})",
            }
        )
        rows.append(row)
    total_row = last_row + 1
    rows.append(["", "Total"] + [None] * (len(columns) - 2))
    for v_index in range(len(table)):
        col = _letter(5 + v_index * 3 + 2)
        formulas.append(
            {
                "cell": f"{col}{total_row}",
                "formula": f"SUM({col}{first_row}:{col}{last_row})",
                "bold": True,
            }
        )
    comparison = {
        "name": "Comparison",
        "columns": columns,
        "rows": rows,
        "formulas": formulas,
        "highlights": highlights,
        "bold_rows": [total_row],
        "number_formats": {"Qty": "0.###"},
    }

    summary_rows = []
    summary_highlights = []
    for position, entry in enumerate(ranking, start=2):
        diff = None if best is None or not entry["complete"] else entry["total"] - best
        pct = (
            None
            if diff is None or not best
            else float(money.to_paise(diff * 100 / best)) / 100
        )
        summary_rows.append(
            [
                entry["rank"],
                entry["vendor"],
                entry["gstin"],
                entry["items_quoted"],
                entry["total"],
                diff,
                pct,
                "" if entry["complete"] else "Did not quote every item",
            ]
        )
        if position == 2 and entry["complete"]:
            summary_highlights += [f"{_letter(c)}{position}" for c in range(1, 9)]
    summary = {
        "name": "Summary",
        "columns": [
            "Rank",
            "Vendor",
            "GSTIN",
            "Items quoted",
            {"name": "Total landed (₹)", "type": "money"},
            {"name": "Above L1 (₹)", "type": "money"},
            {"name": "Above L1 (%)", "type": "percent"},
            "Note",
        ],
        "rows": summary_rows,
        "highlights": summary_highlights,
    }
    return BidAnalysis(sheets=[summary, comparison], ranking=ranking, lowest=lowest)


__all__ = [
    "INR_FORMAT",
    "LOWEST_FILL",
    "BidAnalysis",
    "SpreadsheetError",
    "build_workbook",
    "cost_bid_analysis",
]
