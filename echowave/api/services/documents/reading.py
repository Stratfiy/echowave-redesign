"""Reading an uploaded document back: its text, and its tables as rows.

A vendor's quotation arrives as a PDF, a Word file or a spreadsheet, and the
agent needs the rates in it -- as rows, not as a paragraph of run-together
numbers. Word tables and spreadsheet sheets come back as lists of rows;
PDFs come back as text through the same reader the knowledge base uses,
including its OCR for a scan. Long documents are cut, and say they were.
"""

from __future__ import annotations

import io
import os
import tempfile
from typing import Any

from api.services.documents import sources, templates

MAX_TEXT_CHARS = 20_000
MAX_TABLES = 10
MAX_ROWS = 200
MAX_CELL_CHARS = 200


def _cell(value: Any) -> str:
    text = "" if value is None else str(value)
    return text[:MAX_CELL_CHARS]


def _clip_tables(
    tables: list[list[list[Any]]], notes: list[str]
) -> list[list[list[str]]]:
    if len(tables) > MAX_TABLES:
        notes.append(f"Showing the first {MAX_TABLES} of {len(tables)} tables.")
    out = []
    for index, table in enumerate(tables[:MAX_TABLES], start=1):
        rows = [
            [_cell(c) for c in row]
            for row in table
            if any(c not in (None, "") for c in row)
        ]
        if len(rows) > MAX_ROWS:
            notes.append(
                f"Table {index}: showing the first {MAX_ROWS} of {len(rows)} rows."
            )
            rows = rows[:MAX_ROWS]
        out.append(rows)
    return out


def _xlsx(data: bytes) -> tuple[str, list[list[list[Any]]], list[str]]:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    tables, names = [], []
    for sheet in book.worksheets:
        names.append(sheet.title)
        tables.append([list(row) for row in sheet.iter_rows(values_only=True)])
    return "Sheets: " + ", ".join(names), tables, names


def extract(data: bytes, filename: str) -> dict[str, Any]:
    """Text and tables from a file's bytes, by its extension."""
    from api.services.knowledge_base import extraction

    extension = os.path.splitext(filename)[1].lower()
    notes: list[str] = []
    tables: list[list[list[Any]]] = []
    sheets: list[str] | None = None
    if extension == ".docx":
        paragraphs, tables = templates.text_of(data)
        text = "\n".join(p for p in paragraphs if p.strip())
    elif extension in (".xlsx", ".xlsm"):
        text, tables, sheets = _xlsx(data)
    else:
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, f"document{extension or '.txt'}")
            with open(path, "wb") as handle:
                handle.write(data)
            document = extraction.extract_document(path, filename)
        text = document.full_text
        if document.metadata.get("extractor") == "tesseract":
            notes.append(
                "This PDF is a scan; its text was read by OCR and may have errors."
            )
        if extension == ".pdf":
            notes.append("A PDF's tables are read as text, line by line.")
    if len(text) > MAX_TEXT_CHARS:
        notes.append(
            f"Text cut at {MAX_TEXT_CHARS:,} of {len(text):,} characters; ask for "
            "the part you need."
        )
        text = text[:MAX_TEXT_CHARS]
    out: dict[str, Any] = {
        "filename": filename,
        "text": text,
        "tables": _clip_tables(tables, notes),
    }
    if sheets is not None:
        out["sheets"] = sheets
    if notes:
        out["note"] = " ".join(notes)
    return out


async def find_document(organization_id: int, document: str) -> Any:
    """This workspace's upload by uuid, uuid prefix or file name."""
    from api.db import db_client

    text = (document or "").strip()
    if not text:
        raise sources.SourceError("Say which document: its id or its file name.")
    row = await db_client.get_document_by_uuid(text, organization_id)
    if row is None and len(text) >= 8:
        row = await db_client.find_document_by_uuid_prefix(
            text, organization_id=organization_id
        )
    if row is None:
        wanted = text.lower()
        rows = await db_client.get_documents_for_organization(
            organization_id, limit=200
        )
        exact = [r for r in rows if str(r.filename or "").lower() == wanted]
        partial = [r for r in rows if wanted in str(r.filename or "").lower()]
        matches = exact or partial
        if len(matches) > 1 and not exact:
            names = ", ".join(str(r.filename) for r in matches[:5])
            raise sources.SourceError(
                f"More than one document matches: {names}. Say which."
            )
        row = matches[0] if matches else None
    if row is None:
        raise sources.SourceError(
            "No uploaded document by that id or name in this workspace."
        )
    return row


async def read(organization_id: int, document: str) -> dict[str, Any]:
    from api.services.knowledge_base.errors import KnowledgeBaseError

    row = await find_document(organization_id, document)
    data, filename, _ = await sources.uploaded_bytes(
        organization_id, str(row.document_uuid)
    )
    try:
        result = extract(data, filename)
    except KnowledgeBaseError as exc:
        raise sources.SourceError(
            getattr(exc, "user_message", None) or f"{filename} could not be read."
        ) from exc
    return {"document_uuid": str(row.document_uuid), **result}


__all__ = ["MAX_TEXT_CHARS", "extract", "find_document", "read"]
