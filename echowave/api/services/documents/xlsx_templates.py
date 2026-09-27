"""An Excel workbook as a document template: fields in cells, formulas kept.

The Word engine's convention carries over unchanged -- ``{{field}}`` in a
cell is a field, ``{{items.col}}`` marks an item row -- and one rule is
added for what a workbook has that a Word file does not. **A formula cell
is never written.** The founder's export invoice computes every line amount,
the subtotal, the tax and the total from the cells around them; a fill that
wrote a number over ``=SUM(F18:F23)`` would hand back a sheet that looks
right today and is wrong the first time somebody edits a rate.

Item rows follow from that. They are the template's own pre-made rows,
filled downwards and blanked when unused, never inserted: inserting a row
shifts nothing in openpyxl's formulas, so the ``SUM`` under the items would
quietly stop at the old last row. More items than rows is refused by the
numbers, and the person is told to add rows to the template.

Values keep their type. A quantity goes in as a number, a rate as a number,
so the sheet's arithmetic runs on them; only text stays text.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Iterable
from decimal import Decimal
from typing import Any

import openpyxl

from api.services.documents.placeholders import ITEMS_PREFIX, PLACEHOLDER, TemplateError

MAX_CELLS = 200_000


def is_workbook(data: bytes) -> bool:
    """A zip carrying ``xl/workbook.xml`` is an .xlsx; anything else is not."""
    if data[:2] != b"PK":
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return "xl/workbook.xml" in archive.namelist()
    except zipfile.BadZipFile:
        return False


def _open(data: bytes):
    try:
        return openpyxl.load_workbook(io.BytesIO(data))
    except Exception as exc:
        raise TemplateError(
            "That template could not be opened as an Excel (.xlsx) file."
        ) from exc


def _save(workbook) -> bytes:
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def _is_formula(cell) -> bool:
    return isinstance(cell.value, str) and cell.value.startswith("=")


def _cells(workbook):
    """Every cell with a value, sheet by sheet, in reading order."""
    seen = 0
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                seen += 1
                if seen > MAX_CELLS:
                    raise TemplateError(
                        "That workbook is too large to use as a template."
                    )
                yield sheet, cell


def _item_rows(workbook) -> list[tuple[Any, int]]:
    """(sheet, row number) for each row carrying an ``items.`` placeholder."""
    rows: list[tuple[Any, int]] = []
    for sheet, cell in _cells(workbook):
        if _is_formula(cell) or not isinstance(cell.value, str):
            continue
        if ITEMS_PREFIX in cell.value and PLACEHOLDER.search(cell.value):
            key = (sheet, cell.row)
            if key not in rows:
                rows.append(key)
    return rows


def inspect(data: bytes) -> dict[str, Any]:
    """The fields and item columns the workbook asks for, in reading order,
    and how many item rows it has room for."""
    workbook = _open(data)
    fields: list[str] = []
    columns: list[str] = []
    has_formulas = False
    for _, cell in _cells(workbook):
        if _is_formula(cell):
            has_formulas = True
            continue
        if not isinstance(cell.value, str):
            continue
        for match in PLACEHOLDER.finditer(cell.value):
            name = match.group(1)
            if name.startswith(ITEMS_PREFIX):
                column = name[len(ITEMS_PREFIX) :]
                if column not in columns:
                    columns.append(column)
            elif "." not in name and name not in fields:
                fields.append(name)
    return {
        "fields": fields,
        "item_columns": columns,
        "item_rows": len(_item_rows(workbook)),
        "has_formulas": has_formulas,
    }


def _typed(value: Any) -> Any:
    """What goes in the cell: numbers as numbers, everything else as text."""
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, Decimal):
        return float(value)
    text = str(value)
    stripped = text.replace(",", "").strip().lstrip("₹$€£")
    try:
        number = Decimal(stripped)
    except Exception:  # noqa: BLE001 - not a number, and that is fine
        return text
    if stripped.lstrip("-").replace(".", "", 1).isdigit():
        return (
            int(number)
            if number == number.to_integral_value() and "." not in stripped
            else float(number)
        )
    return text


def _write(cell, resolve) -> None:
    """Replace the placeholders in one cell. A cell that is exactly one
    placeholder takes the value's own type; a cell with words around it
    stays text. An unknown placeholder is left visible."""
    text = cell.value
    whole = PLACEHOLDER.fullmatch(text.strip())
    if whole:
        value = resolve(whole.group(1))
        if value is not None:
            cell.value = _typed(value)
        return

    def sub(match):
        value = resolve(match.group(1))
        return match.group(0) if value is None else str(value)

    cell.value = PLACEHOLDER.sub(sub, text)


def fill(data: bytes, values: dict[str, Any], items: Iterable[dict[str, Any]]) -> bytes:
    """The workbook with every known placeholder replaced, item rows filled
    in place and the rest blanked, and no formula touched."""
    workbook = _open(data)
    items = list(items or [])
    rows = _item_rows(workbook)
    if len(items) > len(rows):
        raise TemplateError(
            f"{len(items)} items were given but the template has room for "
            f"{len(rows)}. Add item rows to the template or split the document."
        )

    def field(name: str):
        if name.startswith(ITEMS_PREFIX):
            return None
        return values.get(name)

    for index, (sheet, row_number) in enumerate(rows):
        item = items[index] if index < len(items) else None
        for cell in sheet[row_number]:
            if (
                cell.value is None
                or _is_formula(cell)
                or not isinstance(cell.value, str)
            ):
                continue
            if ITEMS_PREFIX not in cell.value:
                continue
            if item is None:
                cell.value = None
                continue

            def resolve(name: str, item: dict[str, Any] = item):
                if name.startswith(ITEMS_PREFIX):
                    value = item.get(name[len(ITEMS_PREFIX) :])
                    return "" if value is None else value
                return field(name)

            _write(cell, resolve)
    for _, cell in _cells(workbook):
        if _is_formula(cell) or not isinstance(cell.value, str):
            continue
        if PLACEHOLDER.search(cell.value):
            _write(cell, field)
    return _save(workbook)


def rows_as_text(data: bytes) -> list[tuple[str, list[list[str]]]]:
    """Each sheet's used range as text rows, for a plain rendering. Formula
    cells show their formula: the workbook has never been calculated here."""
    workbook = _open(data)
    sheets: list[tuple[str, list[list[str]]]] = []
    for sheet in workbook.worksheets:
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                rows.append(cells)
        if rows:
            sheets.append((sheet.title, rows))
    return sheets
