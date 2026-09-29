"""The procurement maturity assessment (KAN-182, E-2): seven dimensions,
each scored one to five with a written reason, and what follows from it.

The assessor agent interviews the procurement head and reads a sample of
their documents; this module is where the numbers happen. The model never
adds up a score or picks a band: it hands the seven scores and reasons to
:func:`assess`, which validates them, works out the overall score and its
band, names the gaps, and lays out a 90-day roadmap pointing at the
Decibyl agents that close each one. The report and the score sheet are
built from that result, so the file says what the numbers say.

The seven dimensions follow the maturity-assessment framework procurement
heads are trained on. The framework is a reference; no third party's text
appears here.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

MIN_SCORE, MAX_SCORE = 1, 5
#: A reason shorter than this is a word, not a reason.
MIN_REASON_CHARS = 20
#: A score at or below this is a gap; the next one up is partial.
GAP_AT, PARTIAL_AT = 2, 3


@dataclass(frozen=True)
class Dimension:
    key: str
    name: str
    #: What the interview establishes, in the assessor's words to the person.
    asks: str
    #: The Decibyl agents that close a gap here, in the order to hire them.
    agents: tuple[str, ...]


DIMENSIONS: tuple[Dimension, ...] = (
    Dimension(
        "organisation",
        "Organisation",
        "Who buys, who approves, and whether procurement is a function or a task people do on the side.",
        ("Purchase request desk", "Procurement document drafter"),
    ),
    Dimension(
        "policy",
        "Policy and governance",
        "A written purchase policy, approval limits by value, and whether they are followed.",
        ("Approval matrix", "Audit log"),
    ),
    Dimension(
        "sourcing",
        "Sourcing process",
        "How vendors are found and compared: quotations per order, a comparative statement, who awards.",
        ("RFQ and quotation comparison", "Procurement document drafter"),
    ),
    Dimension(
        "systems",
        "Tools and systems",
        "Where orders, receipts and invoices live: paper, spreadsheets, an ERP; whether they are linked.",
        ("Purchase order follow-up", "Invoice three-way matching"),
    ),
    Dimension(
        "suppliers",
        "Supplier management",
        "How vendors are onboarded and checked (GSTIN, PAN, MSME), and whether there is a vendor master.",
        ("Supplier onboarding executive", "Purchase order follow-up"),
    ),
    Dimension(
        "performance",
        "Performance measurement",
        "Whether savings, on-time delivery and invoice accuracy are measured, and who sees the numbers.",
        ("Spend classifier", "Invoice three-way matching"),
    ),
    Dimension(
        "talent",
        "Talent and skills",
        "Who does the work, what they are trained in, and what is done by hand that need not be.",
        ("Procurement document drafter", "Purchase request desk"),
    ),
)
KEYS = tuple(d.key for d in DIMENSIONS)
BY_KEY = {d.key: d for d in DIMENSIONS}

#: The overall band a mean score falls in: the lower bound of each.
BANDS: tuple[tuple[float, str, str], ...] = (
    (
        4.5,
        "Optimised",
        "Procurement runs on measured, improving processes; the work left is scale.",
    ),
    (
        3.5,
        "Managed",
        "Processes are defined and mostly followed; measurement and systems are the gaps.",
    ),
    (
        2.5,
        "Defined",
        "The basics are written down; following them depends on the person.",
    ),
    (
        1.5,
        "Reactive",
        "Buying happens when something is needed; policy and comparison are occasional.",
    ),
    (0.0, "Ad hoc", "No written policy or process; every purchase is its own story."),
)


class AssessmentError(ValueError):
    """The scores cannot be assessed; the message says what is wrong."""


@dataclass(frozen=True)
class Scored:
    key: str
    name: str
    score: int
    reason: str
    evidence: str


@dataclass(frozen=True)
class Wave:
    days: str
    dimension: str
    score: int
    agents: tuple[str, ...]


@dataclass(frozen=True)
class Assessment:
    client: str
    scored: tuple[Scored, ...]
    overall: float
    band: str
    band_note: str
    gaps: tuple[str, ...]
    partial: tuple[str, ...]
    roadmap: tuple[Wave, ...]
    warnings: tuple[str, ...] = field(default_factory=tuple)

    def as_dict(self) -> dict[str, Any]:
        return {
            "client": self.client,
            "overall": self.overall,
            "band": self.band,
            "band_note": self.band_note,
            "scores": [
                {
                    "dimension": s.name,
                    "key": s.key,
                    "score": s.score,
                    "reason": s.reason,
                }
                for s in self.scored
            ],
            "gaps": list(self.gaps),
            "partial": list(self.partial),
            "roadmap": [
                {
                    "days": w.days,
                    "dimension": w.dimension,
                    "score": w.score,
                    "agents": list(w.agents),
                }
                for w in self.roadmap
            ],
            "warnings": list(self.warnings),
        }


def band_of(overall: float) -> tuple[str, str]:
    for floor, name, note in BANDS:
        if overall >= floor:
            return name, note
    return BANDS[-1][1], BANDS[-1][2]


def _score(raw: Any, key: str) -> int:
    try:
        value = int(raw)
    except (TypeError, ValueError):
        raise AssessmentError(
            f"{key}: score must be a whole number from 1 to 5."
        ) from None
    if not MIN_SCORE <= value <= MAX_SCORE:
        raise AssessmentError(f"{key}: score must be from 1 to 5, not {value}.")
    return value


def assess(client: Any, answers: Any) -> Assessment:
    """Seven scores and reasons in, the assessment out. Raises
    AssessmentError naming the first thing wrong, so the agent asks for it."""
    name = str(client or "").strip()
    if not name:
        raise AssessmentError("Give the client's name for the report.")
    if not isinstance(answers, dict):
        raise AssessmentError("answers must be an object keyed by dimension.")
    given = {str(k).strip().lower(): v for k, v in answers.items()}
    missing = [d.name for d in DIMENSIONS if d.key not in given]
    if missing:
        raise AssessmentError(
            "Every dimension needs a score and a reason; missing: "
            + ", ".join(missing)
            + "."
        )
    unknown = sorted(set(given) - set(KEYS))
    if unknown:
        raise AssessmentError(
            f"Unknown dimension(s) {', '.join(unknown)}; use {', '.join(KEYS)}."
        )

    scored: list[Scored] = []
    warnings: list[str] = []
    for dim in DIMENSIONS:
        entry = given[dim.key]
        if not isinstance(entry, dict):
            raise AssessmentError(f"{dim.key}: give an object with score and reason.")
        score = _score(entry.get("score"), dim.key)
        reason = " ".join(str(entry.get("reason") or "").split())
        if len(reason) < MIN_REASON_CHARS:
            raise AssessmentError(
                f"{dim.key}: the reason must say, in at least {MIN_REASON_CHARS} characters, what was seen or heard."
            )
        evidence = " ".join(str(entry.get("evidence") or "").split())
        if score >= 4 and not evidence:
            warnings.append(
                f"{dim.name} scored {score} on the interview alone; no document was read to confirm it."
            )
        scored.append(Scored(dim.key, dim.name, score, reason[:600], evidence[:600]))

    overall = round(sum(s.score for s in scored) / len(scored), 1)
    band, note = band_of(overall)
    gaps = tuple(s.name for s in scored if s.score <= GAP_AT)
    partial = tuple(s.name for s in scored if s.score == PARTIAL_AT)

    # The roadmap: lowest scores first, ties in dimension order, three waves
    # of thirty days. Everything at or above 4 needs no wave.
    ordered = sorted(
        (s for s in scored if s.score < 4), key=lambda s: (s.score, KEYS.index(s.key))
    )
    labels = ("Days 1-30", "Days 31-60", "Days 61-90")
    waves: list[Wave] = []
    for index, s in enumerate(ordered):
        waves.append(Wave(labels[min(index, 2)], s.name, s.score, BY_KEY[s.key].agents))
    return Assessment(
        client=name,
        scored=tuple(scored),
        overall=overall,
        band=band,
        band_note=note,
        gaps=gaps,
        partial=partial,
        roadmap=tuple(waves),
        warnings=tuple(warnings),
    )


# ---------------------------------------------------------------------------
# The files


def report_docx(result: Assessment, *, assessed_by: str = "", on: str = "") -> bytes:
    """The maturity report as a Word document: one page of findings, the
    score table, a gap diagnosis and the roadmap. Plain python-docx; the
    numbers come from ``result`` and nowhere else."""
    import docx
    from docx.shared import Pt

    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    document.add_heading(f"Procurement maturity assessment: {result.client}", level=1)
    line = f"Overall {result.overall} of 5, {result.band}."
    if on:
        line += f" Assessed {on}."
    if assessed_by:
        line += f" Prepared by {assessed_by}."
    document.add_paragraph(line)
    document.add_paragraph(result.band_note)

    document.add_heading("Scores by dimension", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Dimension", "Score (1-5)", "Why")):
        cell.text = text
    for s in result.scored:
        row = table.add_row().cells
        row[0].text = s.name
        row[1].text = str(s.score)
        row[2].text = s.reason + (f" Evidence: {s.evidence}" if s.evidence else "")

    document.add_heading("Gap diagnosis", level=2)
    if result.gaps:
        document.add_paragraph("Gaps (scored 1 or 2): " + ", ".join(result.gaps) + ".")
    if result.partial:
        document.add_paragraph(
            "Partly in place (scored 3): " + ", ".join(result.partial) + "."
        )
    if not result.gaps and not result.partial:
        document.add_paragraph("No dimension scored below 4.")
    for warning in result.warnings:
        document.add_paragraph(warning)

    document.add_heading("90-day roadmap", level=2)
    if result.roadmap:
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, text in zip(
            table.rows[0].cells,
            ("When", "Dimension", "From", "Agents that close the gap"),
        ):
            cell.text = text
        for w in result.roadmap:
            row = table.add_row().cells
            row[0].text = w.days
            row[1].text = w.dimension
            row[2].text = f"{w.score} of 5"
            row[3].text = ", ".join(w.agents)
    else:
        document.add_paragraph("Nothing to schedule: every dimension is at 4 or above.")

    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def score_sheets(result: Assessment) -> list[dict[str, Any]]:
    """The two sheets of the Excel score sheet, in build_workbook's shape."""
    return [
        {
            "name": "Scores",
            "columns": [
                "Dimension",
                {"name": "Score (1-5)", "type": "number"},
                "Why",
                "Evidence",
            ],
            "rows": [[s.name, s.score, s.reason, s.evidence] for s in result.scored]
            + [["Overall", result.overall, result.band, ""]],
        },
        {
            "name": "Roadmap",
            "columns": [
                "When",
                "Dimension",
                {"name": "Score", "type": "number"},
                "Agents",
            ],
            "rows": [
                [w.days, w.dimension, w.score, ", ".join(w.agents)]
                for w in result.roadmap
            ],
        },
    ]


__all__ = [
    "BANDS",
    "DIMENSIONS",
    "KEYS",
    "Assessment",
    "AssessmentError",
    "assess",
    "band_of",
    "report_docx",
    "score_sheets",
]
